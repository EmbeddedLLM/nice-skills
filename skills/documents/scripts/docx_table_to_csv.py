#!/usr/bin/env python3
"""Pull one table out of a DOCX and write it as CSV.

Reports move data between spreadsheets and Word often enough that a tiny,
predictable exporter is worth having. It reads cell text and nothing else.

    python scripts/docx_table_to_csv.py report.docx --table_index 0 --out table0.csv

Two behaviours worth knowing:

* a merged cell is written once per grid position it occupies, because that is
  how the table grid exposes it;
* a cell holding several paragraphs becomes one CSV field with embedded
  newlines.
"""

from __future__ import annotations

import argparse
import csv
import os
import sys
from pathlib import Path

from docx import Document


def cell_text(cell) -> str:
    """Join a cell's non-empty paragraph texts with newlines."""
    lines = []
    for paragraph in cell.paragraphs:
        text = (paragraph.text or "").strip("\n")
        if text:
            lines.append(text)
    return "\n".join(lines)


def export_table(document: Document, table_index: int, destination: Path) -> None:
    tables = document.tables
    if not tables:
        raise RuntimeError("No tables found in DOCX")
    if not 0 <= table_index < len(tables):
        raise RuntimeError(f"table_index out of range (0..{len(tables) - 1})")

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        for row in tables[table_index].rows:
            writer.writerow([cell_text(cell) for cell in row.cells])


def main() -> None:
    ap = argparse.ArgumentParser(description="Export a DOCX table to CSV")
    ap.add_argument("docx", type=Path)
    ap.add_argument("--table_index", type=int, default=0)
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()

    if not args.docx.exists():
        raise FileNotFoundError(args.docx)

    export_table(Document(str(args.docx)), args.table_index, args.out)
    print(f"[OK] Exported table {args.table_index} → {args.out}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:  # noqa: BLE001 - the CLI reports and exits
        if os.environ.get("DOCS_DEBUG") == "1":
            raise
        print(f"[ERROR] {exc}", file=sys.stderr)
        raise SystemExit(2)
