#!/usr/bin/env python3
"""Report formatting inconsistencies in a DOCX without changing anything.

Answers the "why does this paragraph look different from the others" question
by counting four things:

* runs carrying direct formatting (bold, italic, underline, font, size, colour);
* paragraphs carrying direct formatting (indents, spacing, line spacing);
* paragraphs that read like headings but use no Heading style — the usual
  reason a TOC comes out empty and a screen reader finds no outline;
* which fonts are actually in use, weighted by characters.

Conservative by design: read-only. Cleanup is `style_normalize.py`'s job.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document

EXAMPLE_LIMIT = 5
HEADING_LIKE_LIMIT = 20
HEADING_LIKE_MAX_CHARS = 80
TEXT_PREVIEW = 120
RUN_PREVIEW = 80

REPORT_NOTES = [
    "Direct formatting is not always wrong, but it often causes inconsistent output when templates change.",
    "Heading-like paragraphs not using Heading styles can break TOC and accessibility.",
]


def every_paragraph(document: Document):
    """Body, then table cells, then each section's header and footer."""
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs
    for section in document.sections:
        yield from section.header.paragraphs
        yield from section.footer.paragraphs


def run_has_direct_format(run) -> bool:
    font = run.font
    return any(value is not None for value in (
        run.bold,
        run.italic,
        run.underline,
        font.name,
        font.size,
        font.color.rgb if font.color else None,
    ))


def paragraph_has_direct_format(paragraph) -> bool:
    fmt = paragraph.paragraph_format
    return any(value is not None for value in (
        fmt.left_indent,
        fmt.right_indent,
        fmt.first_line_indent,
        fmt.space_before,
        fmt.space_after,
        fmt.line_spacing,
    ))


def looks_like_heading(paragraph) -> bool:
    """Short, unpunctuated, and carrying a bold run: heading-shaped prose."""
    text = (paragraph.text or "").strip()
    if not text or len(text) > HEADING_LIKE_MAX_CHARS:
        return False
    if text.endswith((".", ";", ":")):
        return False
    return any(run.bold for run in paragraph.runs if run.text and run.bold)


@dataclass
class StyleFindings:
    fonts: Counter = field(default_factory=Counter)
    direct_runs: int = 0
    direct_paragraphs: int = 0
    heading_like: list = field(default_factory=list)
    examples: dict = field(default_factory=lambda: defaultdict(list))

    def add_example(self, kind: str, entry: dict) -> None:
        if len(self.examples[kind]) < EXAMPLE_LIMIT:
            self.examples[kind].append(entry)


def inspect(document: Document) -> StyleFindings:
    findings = StyleFindings()

    for index, paragraph in enumerate(every_paragraph(document), start=1):
        style_name = paragraph.style.name if paragraph.style is not None else ""

        if paragraph_has_direct_format(paragraph):
            findings.direct_paragraphs += 1
            findings.add_example("direct_paragraph_formatting", {
                "para_index": index,
                "text": paragraph.text[:TEXT_PREVIEW],
                "style": style_name,
            })

        if looks_like_heading(paragraph) and not style_name.startswith("Heading"):
            findings.heading_like.append({
                "para_index": index,
                "text": paragraph.text[:TEXT_PREVIEW],
                "style": style_name,
            })

        for run in paragraph.runs:
            if not run.text:
                continue
            if run.font and run.font.name:
                findings.fonts[run.font.name] += len(run.text)
            if run_has_direct_format(run):
                findings.direct_runs += 1
                findings.add_example("direct_run_formatting", {
                    "para_index": index,
                    "run_text": run.text[:RUN_PREVIEW],
                    "style": style_name,
                })

    return findings


def build_report(source: str, findings: StyleFindings) -> dict:
    return {
        "input": str(Path(source).resolve()),
        "fonts_by_char_count": dict(findings.fonts.most_common()),
        "direct_run_formatting_runs": findings.direct_runs,
        "direct_paragraph_formatting_paragraphs": findings.direct_paragraphs,
        "heading_like_paragraphs_not_heading_style": findings.heading_like[:HEADING_LIKE_LIMIT],
        "examples": dict(findings.examples),
        "notes": list(REPORT_NOTES),
    }


def print_summary(findings: StyleFindings) -> None:
    print("[style_lint] direct run-formatting runs:", findings.direct_runs)
    print("[style_lint] direct paragraph-formatting paragraphs:", findings.direct_paragraphs)
    if findings.fonts:
        top = ", ".join(f"{name}({count})" for name, count in findings.fonts.most_common(5))
        print("[style_lint] top fonts by char count:", top)
    if findings.heading_like:
        print("[style_lint] heading-like paragraphs not using Heading styles (first 10):")
        for item in findings.heading_like[:10]:
            print(f"  - #{item['para_index']}: style='{item['style']}' text='{item['text']}'")


def main() -> None:
    ap = argparse.ArgumentParser(description="Lint a DOCX for common style/formatting issues")
    ap.add_argument("input_docx")
    ap.add_argument("--json", dest="json_out", default=None,
                    help="Write a JSON report to this path")
    args = ap.parse_args()

    findings = inspect(Document(args.input_docx))
    report = build_report(args.input_docx, findings)

    if args.json_out:
        Path(args.json_out).write_text(json.dumps(report, indent=2), encoding="utf-8")

    print_summary(findings)


if __name__ == "__main__":
    main()
