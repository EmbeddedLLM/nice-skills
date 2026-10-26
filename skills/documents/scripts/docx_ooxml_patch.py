#!/usr/bin/env python3
"""Low-level DOCX patching for the features python-docx does not cover.

Four jobs, each usable from the CLI or by importing this module:

* switch on Track Revisions in `word/settings.xml`;
* retire an existing `<w:ins>` as a `<w:del>` and add a fresh `<w:ins>` beside
  it, allocating ids that cannot collide with the document's existing ones;
* add a genuine Word comment — body part, anchors, relationship, and content
  type override — anchored by indentation, by a text substring, or by default to
  the first non-empty paragraph;
* three python-docx conveniences: a right-aligned header line, a centred PAGE
  field in the footer, and turning the first paragraph into a hyperlink.

Deliberately pragmatic: these cover common automation tasks, not the whole of
OOXML.

`unzip_docx` / `zip_docx` are the package helpers the sibling scripts in this
directory import, so their names and signatures are part of the contract.
"""

from __future__ import annotations

import argparse
import datetime as dt
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

from lxml import etree

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn as docx_qn
except Exception:
    Document = None  # type: ignore

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CT_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W_NS, "r": R_NS}

DOCUMENT_PART = "word/document.xml"
SETTINGS_PART = "word/settings.xml"
COMMENTS_PART = "word/comments.xml"
DOCUMENT_RELS_PART = "word/_rels/document.xml.rels"
CONTENT_TYPES_PART = "[Content_Types].xml"

COMMENTS_REL_TYPE = f"{R_NS}/comments"
COMMENTS_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"
)
DEFAULT_AUTHOR = "Editor"
PAGE_FIELD_INSTRUCTION = " PAGE \\* MERGEFORMAT "
PAGE_FIELD_CACHED_TEXT = "1"
HYPERLINK_COLOUR = "0000FF"


def _qn(prefix: str, local: str) -> str:
    """Clark-notation name for the handful of namespaces used here."""
    namespaces = {"w": W_NS, "pr": PKG_REL_NS, "ct": CT_NS, "r": R_NS}
    if prefix not in namespaces:
        raise ValueError(prefix)
    return f"{{{namespaces[prefix]}}}{local}"


def iso_now() -> str:
    """Word-style revision timestamp: second precision with a Z suffix."""
    moment = dt.datetime.now(dt.timezone.utc).replace(microsecond=0, tzinfo=None)
    return moment.isoformat() + "Z"


def unzip_docx(docx_path: Path, out_dir: Path) -> None:
    if out_dir.exists():
        shutil.rmtree(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(docx_path, "r") as archive:
        archive.extractall(out_dir)


def zip_docx(in_dir: Path, out_docx_path: Path) -> None:
    if out_docx_path.exists():
        out_docx_path.unlink()
    with zipfile.ZipFile(out_docx_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for root, _dirs, files in os.walk(in_dir):
            for name in files:
                absolute = Path(root) / name
                archive.write(absolute, absolute.relative_to(in_dir).as_posix())


def _parse_part(path: Path) -> etree._ElementTree:
    return etree.parse(str(path), etree.XMLParser(remove_blank_text=False))


def _save_part(tree: etree._ElementTree, path: Path) -> None:
    tree.write(str(path), xml_declaration=True, encoding="UTF-8", standalone="yes")


def _new_part(root: etree._Element, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    etree.ElementTree(root).write(
        str(path), xml_declaration=True, encoding="UTF-8", standalone="yes"
    )


def enable_track_revisions(unzipped: Path) -> None:
    """Make sure settings.xml exists and asks Word to record revisions."""
    settings_path = unzipped / SETTINGS_PART
    if not settings_path.exists():
        _new_part(etree.Element(_qn("w", "settings"), nsmap={"w": W_NS}), settings_path)

    tree = _parse_part(settings_path)
    root = tree.getroot()

    if root.find("w:trackRevisions", namespaces=NS) is None:
        flag = etree.Element(_qn("w", "trackRevisions"))
        compatibility = root.find("w:compat", namespaces=NS)
        if compatibility is not None:
            compatibility.addnext(flag)
        else:
            root.insert(0, flag)

    _save_part(tree, settings_path)


def highest_word_id(root: etree._Element) -> int:
    """Largest integer `w:id` anywhere in the tree — the base for new ones."""
    ids = []
    for element in root.iter():
        value = element.get(_qn("w", "id"))
        if value is None:
            continue
        try:
            ids.append(int(value))
        except ValueError:
            continue
    return max(ids) if ids else 0


def tracked_replace_ins_id(
    unzipped: Path,
    ins_w_id: str,
    new_text: str,
    del_w_id: str = "auto",
    new_ins_w_id: str = "auto",
    author: str = DEFAULT_AUTHOR,
    date_iso: str | None = None,
) -> None:
    """Turn one existing `<w:ins>` into a deletion and add a new insertion."""
    stamp = date_iso or iso_now()
    doc_xml = unzipped / DOCUMENT_PART
    tree = _parse_part(doc_xml)
    root = tree.getroot()

    insertion = root.find(f".//w:ins[@w:id='{ins_w_id}']", namespaces=NS)
    if insertion is None:
        raise RuntimeError(f"Could not find <w:ins w:id='{ins_w_id}'> in document.xml")

    ceiling = highest_word_id(root)
    if del_w_id == "auto":
        del_w_id = str(ceiling + 1)
    if new_ins_w_id == "auto":
        # Stay clear of an explicitly chosen deletion id too.
        try:
            ceiling = max(ceiling, int(del_w_id))
        except ValueError:
            pass
        new_ins_w_id = str(ceiling + 1)

    insertion.tag = _qn("w", "del")
    insertion.attrib[_qn("w", "id")] = del_w_id
    insertion.attrib[_qn("w", "author")] = author
    insertion.attrib[_qn("w", "date")] = stamp
    for text_node in insertion.findall(".//w:t", namespaces=NS):
        text_node.tag = _qn("w", "delText")  # deleted text is never w:t

    parent = insertion.getparent()
    position = parent.index(insertion)

    new_insertion = etree.Element(_qn("w", "ins"))
    new_insertion.attrib[_qn("w", "id")] = new_ins_w_id
    new_insertion.attrib[_qn("w", "author")] = author
    new_insertion.attrib[_qn("w", "date")] = stamp
    run = etree.SubElement(new_insertion, _qn("w", "r"))
    etree.SubElement(run, _qn("w", "t")).text = new_text

    parent.insert(position + 1, new_insertion)
    _save_part(tree, doc_xml)


def paragraph_plain_text(paragraph: etree._Element) -> str:
    return "".join(node.text or "" for node in paragraph.findall(".//w:t", namespaces=NS))


class CommentAnchorSelector:
    """Picks the paragraph a new comment attaches to."""

    def __init__(self, paragraphs: list[etree._Element], indent_left_twips: int | None,
                 contains: str | None) -> None:
        self.paragraphs = paragraphs
        self.indent_left_twips = indent_left_twips
        self.contains = contains

    def select(self) -> etree._Element:
        if self.indent_left_twips is None and self.contains is None:
            for paragraph in self.paragraphs:
                if paragraph_plain_text(paragraph).strip():
                    return paragraph
            return self.paragraphs[0]

        wanted_indent = None if self.indent_left_twips is None else str(self.indent_left_twips)
        for paragraph in self.paragraphs:
            if wanted_indent is not None and not self._indent_matches(paragraph, wanted_indent):
                continue
            if self.contains is not None and self.contains not in paragraph_plain_text(paragraph):
                continue
            return paragraph

        raise RuntimeError(self._no_match_message())

    @staticmethod
    def _indent_matches(paragraph: etree._Element, wanted: str) -> bool:
        properties = paragraph.find("w:pPr", namespaces=NS)
        indent = properties.find("w:ind", namespaces=NS) if properties is not None else None
        return indent is not None and indent.get(_qn("w", "left")) == wanted

    def _no_match_message(self) -> str:
        criteria = []
        if self.indent_left_twips is not None:
            criteria.append(f"indent_left_twips={self.indent_left_twips}")
        if self.contains is not None:
            criteria.append(f"contains={self.contains!r}")
        hint = (
            "Tip: omit both flags to anchor to the first paragraph, "
            "or use --contains to match a short snippet from the target paragraph."
        )
        return ("Could not find paragraph matching predicate ("
                + ", ".join(criteria) + ").\n" + hint)


def next_comment_id(unzipped: Path, document_root: etree._Element) -> str:
    """Smallest unused non-negative id across anchors and existing comments."""
    used: set[int] = set()

    anchors = (document_root.findall(".//w:commentRangeStart", namespaces=NS)
               + document_root.findall(".//w:commentRangeEnd", namespaces=NS))
    comments_path = unzipped / COMMENTS_PART
    if comments_path.exists():
        try:
            anchors += _parse_part(comments_path).getroot().findall("w:comment", namespaces=NS)
        except Exception:
            pass

    for element in anchors:
        value = element.get(_qn("w", "id"))
        try:
            used.add(int(value))
        except (TypeError, ValueError):
            continue

    candidate = 0
    while candidate in used:
        candidate += 1
    return str(candidate)


def ensure_comments_relationship(unzipped: Path) -> None:
    """Add the document → comments.xml relationship when it is missing."""
    rels_path = unzipped / DOCUMENT_RELS_PART
    rels_path.parent.mkdir(parents=True, exist_ok=True)

    if rels_path.exists():
        tree = _parse_part(rels_path)
        root = tree.getroot()
    else:
        root = etree.Element(_qn("pr", "Relationships"), nsmap={"pr": PKG_REL_NS})
        tree = etree.ElementTree(root)

    relationships = root.findall(".//pr:Relationship", namespaces={"pr": PKG_REL_NS})
    if any(rel.get("Type") == COMMENTS_REL_TYPE for rel in relationships):
        return

    highest = 0
    for relationship in relationships:
        identifier = relationship.get("Id", "")
        if identifier.startswith("rId"):
            try:
                highest = max(highest, int(identifier[3:]))
            except ValueError:
                continue

    added = etree.SubElement(root, _qn("pr", "Relationship"))
    added.set("Id", f"rId{highest + 1}")
    added.set("Type", COMMENTS_REL_TYPE)
    added.set("Target", "comments.xml")
    _save_part(tree, rels_path)


def ensure_comments_content_type(unzipped: Path) -> None:
    content_types_path = unzipped / CONTENT_TYPES_PART
    tree = _parse_part(content_types_path)
    root = tree.getroot()

    declared = any(
        override.get("PartName") == f"/{COMMENTS_PART}"
        for override in root.findall(".//ct:Override", namespaces={"ct": CT_NS})
    )
    if declared:
        return

    override = etree.SubElement(root, _qn("ct", "Override"))
    override.set("PartName", f"/{COMMENTS_PART}")
    override.set("ContentType", COMMENTS_CONTENT_TYPE)
    _save_part(tree, content_types_path)


def anchor_comment_in_paragraph(paragraph: etree._Element, comment_id: str) -> None:
    """Place range start/end markers and the reference run around the runs."""
    runs = paragraph.findall("w:r", namespaces=NS)
    if not runs:
        runs = [etree.SubElement(paragraph, _qn("w", "r"))]

    start = etree.Element(_qn("w", "commentRangeStart"))
    start.set(_qn("w", "id"), comment_id)
    paragraph.insert(paragraph.index(runs[0]), start)

    end = etree.Element(_qn("w", "commentRangeEnd"))
    end.set(_qn("w", "id"), comment_id)
    paragraph.insert(paragraph.index(runs[-1]) + 1, end)

    reference_run = etree.Element(_qn("w", "r"))
    etree.SubElement(reference_run, _qn("w", "commentReference")).set(_qn("w", "id"), comment_id)
    paragraph.insert(paragraph.index(end) + 1, reference_run)


def append_comment_entry(unzipped: Path, comment_id: str, text: str, author: str,
                         date_iso: str) -> None:
    comments_path = unzipped / COMMENTS_PART
    tree = _parse_part(comments_path)
    root = tree.getroot()

    if root.find(f".//w:comment[@w:id='{comment_id}']", namespaces=NS) is not None:
        raise RuntimeError(f"comments.xml already contains comment id={comment_id}")

    comment = etree.SubElement(root, _qn("w", "comment"))
    comment.set(_qn("w", "id"), comment_id)
    comment.set(_qn("w", "author"), author)
    comment.set(_qn("w", "date"), date_iso)
    run = etree.SubElement(etree.SubElement(comment, _qn("w", "p")), _qn("w", "r"))
    etree.SubElement(run, _qn("w", "t")).text = text

    _save_part(tree, comments_path)


def add_comment_to_first_indented_paragraph(
    unzipped: Path,
    comment_text: str,
    indent_left_twips: int | None = None,
    contains: str | None = None,
    comment_id: str = "auto",
    author: str = DEFAULT_AUTHOR,
    date_iso: str | None = None,
) -> None:
    """Add a Word comment, wiring up every part that has to agree on its id.

    Anchor selection:

    * with neither predicate, the first non-empty paragraph (falling back to the
      very first paragraph);
    * with `indent_left_twips`, a paragraph whose `w:pPr/w:ind/@w:left` equals
      that value — note that most paragraphs declare no indentation at all;
    * with `contains`, a paragraph whose visible text holds that substring.

    `comment_id="auto"` picks the next id that collides with neither an existing
    anchor nor an existing comment. The unzipped package is edited in place.
    """
    stamp = date_iso or iso_now()

    doc_xml = unzipped / DOCUMENT_PART
    tree = _parse_part(doc_xml)
    root = tree.getroot()

    paragraphs = list(root.findall(".//w:p", namespaces=NS))
    if not paragraphs:
        raise RuntimeError("No paragraphs found in document.xml; cannot anchor comment.")

    target = CommentAnchorSelector(paragraphs, indent_left_twips, contains).select()

    if comment_id == "auto":
        comment_id = next_comment_id(unzipped, root)

    comments_path = unzipped / COMMENTS_PART
    if not comments_path.exists():
        _new_part(etree.Element(_qn("w", "comments"), nsmap={"w": W_NS}), comments_path)

    ensure_comments_relationship(unzipped)
    ensure_comments_content_type(unzipped)

    anchor_comment_in_paragraph(target, comment_id)
    _save_part(tree, doc_xml)

    append_comment_entry(unzipped, comment_id, comment_text, author, stamp)


def require_python_docx() -> None:
    if Document is None:
        raise RuntimeError("python-docx is not available in this environment")


def add_header_date(docx_path: Path, date_text: str) -> None:
    """Right-aligned line in the first section's header."""
    require_python_docx()
    document = Document(str(docx_path))
    header = document.sections[0].header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.text = date_text
    document.save(str(docx_path))


def add_page_number_field_to_footer(docx_path: Path) -> None:
    """Centred PAGE field in the first section's footer."""
    require_python_docx()
    document = Document(str(docx_path))
    paragraph = document.sections[0].footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = paragraph.add_run()
    for element in _page_field_elements():
        run._r.append(element)

    document.save(str(docx_path))


def _page_field_elements() -> list:
    """The five pieces of a complex PAGE field with a cached result of "1"."""
    begin = OxmlElement("w:fldChar")
    begin.set(docx_qn("w:fldCharType"), "begin")

    instruction = OxmlElement("w:instrText")
    instruction.set(docx_qn("xml:space"), "preserve")
    instruction.text = PAGE_FIELD_INSTRUCTION

    separate = OxmlElement("w:fldChar")
    separate.set(docx_qn("w:fldCharType"), "separate")

    cached = OxmlElement("w:t")
    cached.text = PAGE_FIELD_CACHED_TEXT

    end = OxmlElement("w:fldChar")
    end.set(docx_qn("w:fldCharType"), "end")

    return [begin, instruction, separate, cached, end]


def hyperlink_first_paragraph(docx_path: Path, url: str, underline: bool = True) -> None:
    """Replace the first paragraph with one linked run holding the same text.

    Blunt by design: per-run formatting inside that paragraph is dropped. Edge
    whitespace survives via `xml:space="preserve"`.
    """
    require_python_docx()
    document = Document(str(docx_path))
    if not document.paragraphs:
        raise RuntimeError("Document has no paragraphs")

    paragraph = document.paragraphs[0]
    text = paragraph.text or ""
    for run in reversed(paragraph.runs):
        paragraph._p.remove(run._r)

    relationship_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(docx_qn("r:id"), relationship_id)

    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(docx_qn("w:val"), HYPERLINK_COLOUR)
    properties.append(colour)
    if underline:
        marker = OxmlElement("w:u")
        marker.set(docx_qn("w:val"), "single")
        properties.append(marker)
    run.append(properties)

    text_node = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace() or "  " in text:
        text_node.set(docx_qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)

    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    document.save(str(docx_path))


# ---------------------------------------------------------------- CLI passes


def run_ooxml_pass(args: argparse.Namespace) -> None:
    """Unzip, apply the OOXML-level operations, repack.

    The python-docx helpers above may already have rewritten the *output* copy,
    so this pass always works on the current document — `--out` when given,
    otherwise the input.
    """
    current = Path(args.out or args.docx)
    if not current.exists():
        raise FileNotFoundError(current)

    workdir = Path(tempfile.mkdtemp(prefix="docx_patch_"))
    unzipped = workdir / "unz"
    unzip_docx(current, unzipped)

    if args.enable_track:
        enable_track_revisions(unzipped)

    if args.tracked_replace_ins_id is not None:
        tracked_replace_ins_id(
            unzipped=unzipped,
            ins_w_id=args.tracked_replace_ins_id,
            new_text=args.new_text,
            del_w_id=args.del_id,
            new_ins_w_id=args.ins_id,
            author=args.author,
            date_iso=args.date,
        )

    if args.add_comment:
        add_comment_to_first_indented_paragraph(
            unzipped=unzipped,
            comment_text=args.comment_text,
            indent_left_twips=args.indent_left_twips,
            contains=args.contains,
            comment_id=args.comment_id,
            author=args.author,
            date_iso=args.date,
        )

    zip_docx(unzipped, current)
    shutil.rmtree(workdir, ignore_errors=True)


def run_python_docx_pass(args: argparse.Namespace, docx_path: Path) -> list[str]:
    """Apply the python-docx helpers, returning the operations performed."""
    operations: list[str] = []

    if args.header_date is None and not args.add_page_numbers and args.hyperlink_first is None:
        return operations

    require_python_docx()
    if args.header_date is not None:
        add_header_date(docx_path, args.header_date)
        operations.append(f"header-date={args.header_date!r}")
    if args.add_page_numbers:
        add_page_number_field_to_footer(docx_path)
        operations.append("add-page-numbers")
    if args.hyperlink_first is not None:
        hyperlink_first_paragraph(docx_path, args.hyperlink_first)
        operations.append(f"hyperlink-first={args.hyperlink_first!r}")

    return operations


def summarise(out_path: Path, operations: list[str]) -> None:
    print(f"[OK] Patched -> {out_path}")
    for operation in operations:
        print(f"  - {operation}")


def build_parser() -> argparse.ArgumentParser:
    ap = argparse.ArgumentParser(description="DOCX patch helper (OOXML + python-docx utilities)")
    ap.add_argument("docx", help="Input DOCX")
    ap.add_argument("--out", default=None, help="Output DOCX (default: overwrite input)")

    ap.add_argument("--enable-track", action="store_true",
                    help="Enable <w:trackRevisions/> in settings.xml")
    ap.add_argument("--author", default=DEFAULT_AUTHOR, help="Author metadata")
    ap.add_argument("--date", default=None, help="ISO date for metadata (default: now)")

    ap.add_argument("--tracked-replace-ins-id", default=None,
                    help="Convert <w:ins w:id=...> into delete+insert")
    ap.add_argument("--new-text", default="", help="New text for tracked insertion")
    ap.add_argument("--del-id", default="auto",
                    help="w:id for the deletion (default: auto-pick a non-colliding id)")
    ap.add_argument("--ins-id", default="auto",
                    help="w:id for the new insertion (default: auto-pick a non-colliding id)")

    ap.add_argument("--add-comment", action="store_true",
                    help="Add a Word comment via comments.xml")
    ap.add_argument("--comment-text", default="", help="Comment body")
    ap.add_argument("--indent-left-twips", type=int, default=None,
                    help="Match paragraphs by indentation left (twips)")
    ap.add_argument("--contains", default=None,
                    help="Require substring in paragraph text")
    ap.add_argument("--comment-id", default="auto",
                    help="Comment id (default: auto-pick a non-colliding id)")

    ap.add_argument("--header-date", default=None,
                    help="Set first section header to this text (right aligned)")
    ap.add_argument("--add-page-numbers", action="store_true",
                    help="Add a centered page number field to footer")
    ap.add_argument("--hyperlink-first", default=None,
                    help="Add hyperlink to first paragraph text to this URL")
    return ap


def main() -> int:
    args = build_parser().parse_args()

    docx_path = Path(args.out or args.docx)
    if args.out:
        shutil.copyfile(args.docx, args.out)

    # python-docx first: those helpers rewrite whole parts, so the finer-grained
    # OOXML pass has to run against what they leave behind.
    operations = run_python_docx_pass(args, docx_path)

    if args.enable_track or args.tracked_replace_ins_id or args.add_comment:
        if args.enable_track:
            operations.append("enable-track")
        if args.tracked_replace_ins_id:
            operations.append(f"tracked-replace-ins-id={args.tracked_replace_ins_id!r}")
        if args.add_comment:
            operations.append("add-comment")
        run_ooxml_pass(args)

    if operations:
        summarise(Path(args.out or args.docx), operations)
    else:
        print("[OK] No changes requested (nothing to do)")

    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # noqa: BLE001 - the CLI reports and exits
        if os.environ.get("DOCS_DEBUG") == "1":
            raise
        print(f"[ERROR] {exc}", file=sys.stderr)
        raise SystemExit(2)
