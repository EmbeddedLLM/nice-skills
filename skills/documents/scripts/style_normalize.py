#!/usr/bin/env python3
"""Reduce formatting drift by handing appearance back to the styles.

The conservative companion to `style_lint.py`. By default it clears *run-level*
overrides — bold, italic, underline, font name, size, colour — so the paragraph
and style definitions decide how text looks again.

Two opt-in passes go further:

* `--clear_paragraph_format` also clears indents and spacing overrides. It can
  move layout; use it when the user wants spacing to come from styles.
* `--enforce_heading_spacing` pins space-after on heading paragraphs to 6pt.

Clearing overrides can change pagination, so render to PNG afterwards and read
the pages.

Both output spellings are accepted; the positional form came first and `--out`
is the alias the rest of this package uses:

    python scripts/style_normalize.py in.docx out.docx
    python scripts/style_normalize.py in.docx --out out.docx
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.shared import Pt

RUN_OVERRIDE_ATTRIBUTES = ("name", "size", "bold", "italic", "underline")
INDENT_ATTRIBUTES = ("left_indent", "right_indent", "first_line_indent")
SPACING_ATTRIBUTES = ("space_before", "space_after", "line_spacing", "line_spacing_rule")
HEADING_SPACE_AFTER_PT = 6.0


def editable_paragraphs(document: Document):
    """Body paragraphs plus every table cell paragraph.

    Headers and footers are left alone: their formatting is usually deliberate.
    """
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def is_heading(style_name: str | None) -> bool:
    return bool(style_name) and style_name.lower().startswith("heading")


def clear_run_overrides(document: Document) -> int:
    """Drop direct run formatting; count the runs that had any."""
    cleared = 0
    for paragraph in editable_paragraphs(document):
        for run in paragraph.runs:
            font = run.font
            overridden = (
                any(getattr(font, attribute) is not None for attribute in RUN_OVERRIDE_ATTRIBUTES)
                or font.color.rgb is not None
            )
            if not overridden:
                continue
            for attribute in RUN_OVERRIDE_ATTRIBUTES:
                setattr(font, attribute, None)
            font.color.rgb = None
            cleared += 1
    return cleared


def clear_paragraph_overrides(document: Document) -> int:
    """Drop indents and spacing overrides; each attribute counts separately."""
    cleared = 0
    for paragraph in editable_paragraphs(document):
        fmt = paragraph.paragraph_format
        for attribute in (*INDENT_ATTRIBUTES, *SPACING_ATTRIBUTES):
            if getattr(fmt, attribute) is not None:
                setattr(fmt, attribute, None)
                cleared += 1
    return cleared


def apply_heading_spacing(document: Document,
                          space_after_pt: float = HEADING_SPACE_AFTER_PT) -> int:
    wanted = Pt(space_after_pt)
    updated = 0
    for paragraph in document.paragraphs:
        style_name = paragraph.style.name if paragraph.style is not None else None
        if not is_heading(style_name):
            continue
        if paragraph.paragraph_format.space_after != wanted:
            paragraph.paragraph_format.space_after = wanted
            updated += 1
    return updated


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("input_docx")
    ap.add_argument("output_docx", nargs="?", help="Output DOCX path (positional)")
    ap.add_argument("--out", default=None,
                    help="Output DOCX path (alias for positional output_docx)")
    ap.add_argument(
        "--clear_paragraph_format",
        action="store_true",
        help="Also clear paragraph-level direct formatting overrides (more invasive).",
    )
    ap.add_argument(
        "--enforce_heading_spacing",
        action="store_true",
        help="Enforce a simple heading spacing rule (space-after = 6pt).",
    )
    args = ap.parse_args()

    destination = args.out or args.output_docx
    if not destination:
        ap.error("output_docx (positional) or --out is required")

    document = Document(args.input_docx)
    run_changes = clear_run_overrides(document)
    paragraph_changes = clear_paragraph_overrides(document) if args.clear_paragraph_format else 0
    heading_changes = apply_heading_spacing(document) if args.enforce_heading_spacing else 0

    out_path = Path(destination)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out_path))

    print(
        f"[OK] wrote {out_path} | run_overrides_cleared={run_changes} "
        f"para_overrides_cleared={paragraph_changes} heading_spacing_updates={heading_changes}"
    )


if __name__ == "__main__":
    main()
