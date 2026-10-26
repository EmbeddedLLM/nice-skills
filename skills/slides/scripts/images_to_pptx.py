#!/usr/bin/env python3
"""Compose a local image plan into a centered raster deck with speaker notes.

Adapted from DeerFlow ppt-generation; see ../SOURCES.md and ../LICENSE-deer-flow.txt.
"""
import argparse
from io import BytesIO
import json
import os
from pathlib import Path
import tempfile

from PIL import Image, ImageOps
from pptx import Presentation
from pptx.util import Inches


def build(plan_path, image_paths, output, force=False):
    if output.suffix.lower() != '.pptx':
        raise ValueError('output must end in .pptx')
    if output.exists() and not force:
        raise ValueError('output exists; use --force to replace it')
    if output.resolve() in {p.resolve() for p in [plan_path, *image_paths]}:
        raise ValueError('output must differ from every input')
    plan = json.loads(plan_path.read_text())
    ratio = plan.get('aspect_ratio', '16:9')
    if ratio not in ('16:9', '4:3'):
        raise ValueError('aspect_ratio must be 16:9 or 4:3')
    slides = plan['slides']
    if not isinstance(slides, list) or not slides or len(slides) != len(image_paths):
        raise ValueError('plan must contain exactly one slide entry per image')
    presentation = Presentation()
    presentation.slide_height = Inches(7.5)
    presentation.slide_width = Inches(40 / 3 if ratio == '16:9' else 10)
    width, height = presentation.slide_width, presentation.slide_height
    for info, path in zip(slides, image_paths):
        if not isinstance(info, dict):
            raise ValueError('each slide entry must be an object')
        with Image.open(path) as original:
            image = ImageOps.exif_transpose(original).convert('RGBA')
            scale = min(width / image.width, height / image.height)
            draw_width, draw_height = round(image.width * scale), round(image.height * scale)
            data = BytesIO()
            image.save(data, format='PNG')
            data.seek(0)
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(data, (width - draw_width) // 2, (height - draw_height) // 2,
                                 draw_width, draw_height)
        notes = []
        for field in ('title', 'subtitle', 'notes', 'credits'):
            if field in info:
                if not isinstance(info[field], str):
                    raise ValueError(f'{field} must be a string')
                notes.append(info[field])
        points = info.get('key_points', [])
        if not isinstance(points, list) or not all(isinstance(p, str) for p in points):
            raise ValueError('key_points must be a list of strings')
        notes.extend(points)
        slide.notes_slide.notes_text_frame.text = '\n'.join(notes)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.pptx', delete=False) as temporary:
        staged = Path(temporary.name)
    try:
        presentation.save(staged)
        os.replace(staged, output)
    finally:
        staged.unlink(missing_ok=True)
    return {'output': str(output), 'slides': len(slides), 'aspect_ratio': ratio, 'editable_content': False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--plan-file', required=True, type=Path)
    parser.add_argument('--slide-images', required=True, nargs='+', type=Path)
    parser.add_argument('--output-file', required=True, type=Path)
    parser.add_argument('--force', action='store_true')
    args = parser.parse_args()
    try:
        print(json.dumps(build(args.plan_file, args.slide_images, args.output_file, args.force)))
    except (ValueError, KeyError, TypeError, OSError) as error:
        parser.exit(1, f'error: {error}\n')


if __name__ == '__main__':
    main()
