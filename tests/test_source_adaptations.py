"""Exercise portable theme, design-search, raster-deck, and source-package contracts."""
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest

from PIL import Image
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]


def run(script, *args, cwd=None):
    return subprocess.run([sys.executable, str(ROOT / script), *map(str, args)],
                          cwd=cwd, capture_output=True, text=True)


class ThemeTests(unittest.TestCase):
    def test_every_material_exports_real_tokens_and_its_notice(self):
        """Ensures every bundled theme is discoverable and exports usable CSS with its license."""
        script = 'skills/theme-factory/scripts/hallmark.py'
        listed = run(script, 'list')
        self.assertEqual(listed.returncode, 0, listed.stderr)
        entries = json.loads(listed.stdout)
        self.assertEqual(len(entries), 48)
        self.assertEqual(len({r['slug'] for r in entries}), 48)
        with tempfile.TemporaryDirectory() as tmp:
            for row in entries:
                with self.subTest(theme=row['slug']):
                    output = Path(tmp) / (row['slug'] + '.css')
                    result = run(script, 'css', row['slug'], '--output', output, cwd=tmp)
                    self.assertEqual(result.returncode, 0, result.stderr)
                    self.assertIn(':root', output.read_text())
                    self.assertIn('--color-', output.read_text())
                    self.assertEqual(output.with_suffix('.css.LICENSE.txt').read_bytes(),
                                     (ROOT / 'skills/theme-factory/LICENSE-hallmark.txt').read_bytes())
            before = output.read_bytes()
            self.assertNotEqual(run(script, 'css', row['slug'], '--output', output).returncode, 0)
            self.assertEqual(output.read_bytes(), before)
            self.assertNotEqual(run(script, 'show', '../../LICENSE').returncode, 0)
            output.unlink()
            notice = output.with_suffix('.css.LICENSE.txt')
            notice.write_text('Existing notice')
            self.assertNotEqual(run(script, 'css', row['slug'], '--output', output).returncode, 0)
            self.assertFalse(output.exists())
            self.assertEqual(notice.read_text(), 'Existing notice')


class DesignSearchTests(unittest.TestCase):
    def test_upstream_semantic_regressions_in_isolated_process(self):
        """Ensures retrieval, mode, stack, and data regressions without generic module-name collisions."""
        result = subprocess.run([sys.executable, '-m', 'unittest', 'discover', '-s',
                                 str(ROOT / 'tests/design_research'), '-p', 'test_*.py', '-v'],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    def test_catalog_and_task_specific_queries_work_offline(self):
        """Ensures portable search finds relevant guidance and generates a plan without writes."""
        validator = run('skills/design-research/scripts/validate_data.py')
        self.assertEqual(validator.returncode, 0, validator.stderr + validator.stdout)
        with tempfile.TemporaryDirectory() as tmp:
            script = 'skills/design-research/scripts/search.py'
            result = run(script, 'accessible analytics dashboard', '--domain', 'ux', '--max-results', '3', cwd=tmp)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertIn('Accessibility', result.stdout)
            plan = run(script, 'public library booking', '--design-system', '-p', 'Library Booking', cwd=tmp)
            self.assertEqual(plan.returncode, 0, plan.stderr)
            self.assertIn('Library Booking'.upper(), plan.stdout.upper())
            self.assertEqual(list(Path(tmp).iterdir()), [])
            rejected = run(script, 'forms', '--stack', 'nonexistent-stack', cwd=tmp)
            self.assertNotEqual(rejected.returncode, 0)


class ImageDeckTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.work = Path(self.temporary.name)
        self.plan = self.work / 'plan.json'
        self.output = self.work / 'deck.pptx'
        self.images = [self.work / 'wide.png', self.work / 'tall.png']
        Image.new('RGB', (400, 100), 'red').save(self.images[0])
        Image.new('RGBA', (100, 400), (0, 0, 255, 100)).save(self.images[1])
        self.plan.write_text(json.dumps({'aspect_ratio': '16:9', 'slides': [
            {'title': 'Wide', 'credits': 'Synthetic red fixture'}, {'title': 'Tall'}]}))

    def build(self, *extra):
        return run('skills/slides/scripts/images_to_pptx.py', '--plan-file', self.plan,
                   '--slide-images', *self.images, '--output-file', self.output, *extra)

    def test_centered_contain_order_transparency_and_notes(self):
        """Ensures wide and tall images are centered without distortion and keep notes/order."""
        result = self.build()
        self.assertEqual(result.returncode, 0, result.stderr)
        deck = Presentation(self.output)
        self.assertEqual(len(deck.slides), 2)
        for slide, expected_ratio in zip(deck.slides, [4, .25]):
            picture = slide.shapes[0]
            self.assertAlmostEqual(picture.width / picture.height, expected_ratio, places=4)
            self.assertLessEqual(abs(2 * picture.left + picture.width - deck.slide_width), 1)
            self.assertLessEqual(abs(2 * picture.top + picture.height - deck.slide_height), 1)
        self.assertIn('Synthetic red fixture', deck.slides[0].notes_slide.notes_text_frame.text)
        self.assertEqual(deck.slides[1].notes_slide.notes_text_frame.text, 'Tall')
        self.assertEqual(deck.slides[1].shapes[0].image.ext, 'png')

    def test_invalid_plan_or_missing_image_preserves_previous_output(self):
        """Ensures validation and decode failures cannot replace an existing presentation."""
        self.output.write_bytes(b'previous presentation')
        self.assertNotEqual(self.build().returncode, 0)
        for plan in [{'aspect_ratio': 'bad', 'slides': [{}, {}]}, {'slides': [{}]}, {'slides': []}]:
            self.plan.write_text(json.dumps(plan))
            self.assertNotEqual(self.build('--force').returncode, 0)
            self.assertEqual(self.output.read_bytes(), b'previous presentation')
        self.plan.write_text(json.dumps({'slides': [{}, {}]}))
        self.images[1].unlink()
        self.assertNotEqual(self.build('--force').returncode, 0)
        self.assertEqual(self.output.read_bytes(), b'previous presentation')
        self.assertEqual(sorted(p.name for p in self.work.iterdir()), ['deck.pptx', 'plan.json', 'wide.png'])


if __name__ == '__main__':
    unittest.main()
