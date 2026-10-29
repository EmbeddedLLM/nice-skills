"""Contracts for local helpers that replace tools in the upstream test corpus."""

import json
import os
import shutil
from pathlib import Path
import sys
import unittest
import zipfile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook

from skill_scripts import SKILLS, ScriptCase, require_binary


class FrontendTests(ScriptCase):
    def test_starter_bundles_offline_and_preserves_existing_outputs(self):
        """A generated starter exports with its code inline and refuses overwrite."""
        scripts = SKILLS / 'frontend-design/scripts'
        project = self.work / 'project'
        self.run_python(scripts / 'init_artifact.py', project)
        output = self.work / 'artifact.html'
        self.run_python(scripts / 'bundle_artifact.py', project / 'index.html', output)
        text = output.read_text()
        self.assertIn("document.querySelector", text)
        self.assertIn('<style>', text)
        self.assertNotIn('src="app.js"', text)
        before = output.read_bytes()
        self.run_python(scripts / 'bundle_artifact.py', project / 'index.html', output, code=1)
        self.assertEqual(output.read_bytes(), before)
        result = self.run_python(scripts / 'init_artifact.py', project, code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertTrue((project / 'app.js').is_file())

    def test_unsupported_or_missing_assets_publish_nothing(self):
        """Static exports refuse inputs they cannot preserve faithfully."""
        for markup in ('<script type="module" src="app.js"></script>',
                       '<img src="missing.png">', '<img src="../outside.png">',
                       '<style>@import "remote.css";</style>'):
            with self.subTest(markup=markup):
                entry = self.write('index.html', markup)
                output = self.work / 'out.html'
                self.run_python(SKILLS / 'frontend-design/scripts/bundle_artifact.py',
                                entry, output, code=1)
                self.assertFalse(output.exists())


class ThemeTests(ScriptCase):
    script_path = SKILLS / 'theme-factory/scripts/themes.py'

    def test_catalog_and_colour_validation(self):
        """The complete theme catalog passes its own contrast and Office XML checks."""
        report = self.run_python(self.script_path, 'verify').json()
        self.assertEqual(report['status'], 'passed')
        rows = self.run_python(self.script_path, 'list', '--json').json()
        self.assertEqual(len(rows), report['themes'])
        self.assertEqual(len({r['id'] for r in rows}), len(rows))
        contrast = self.run_python(self.script_path, 'contrast', '#000000', '#FFFFFF').json()
        self.assertEqual(contrast['ratio'], 21)

    def test_export_is_repeatable_and_preflights_all_outputs(self):
        """Repeated exports preserve bytes; conflicting files prevent partial publication."""
        row = self.run_python(self.script_path, 'list', '--json').json()[0]
        output = self.work / 'theme'
        args = ('export', row['id'], '--out', output)
        self.run_python(self.script_path, *args)
        before = {p.name: p.read_bytes() for p in output.iterdir()}
        self.assertEqual(set(before), {'theme.json', 'theme.css', 'theme.xml', 'SOURCE.txt'})
        self.run_python(self.script_path, *args)
        self.assertEqual(before, {p.name: p.read_bytes() for p in output.iterdir()})
        (output / 'SOURCE.txt').write_text('user content')
        (output / 'theme.css').unlink()
        self.run_python(self.script_path, *args, code=1)
        self.assertFalse((output / 'theme.css').exists())
        self.assertEqual((output / 'SOURCE.txt').read_text(), 'user content')


class SpreadsheetTests(ScriptCase):
    script_path = SKILLS / 'spreadsheets/scripts/recalculate.py'

    def workbook(self):
        book = Workbook()
        book.active['A1'] = 2
        book.active['A2'] = '=A1*3'
        source = self.work / 'book.xlsx'
        book.save(source)
        return source

    def test_check_only_reports_uncached_formulas_without_mutation(self):
        """Uncalculated formulas are unverified, and inspection leaves the file intact."""
        source = self.workbook()
        before = source.read_bytes()
        result = self.run_python(self.script_path, source, '--check-only', code=2).json()
        self.assertEqual(result['status'], 'unverified')
        self.assertEqual(result['missing_caches'], ['Sheet!A2'])
        self.assertFalse(result['recalculated'])
        self.assertEqual(source.read_bytes(), before)

    def test_invalid_options_and_existing_outputs_are_rejected(self):
        """Validation must not replace the input or an existing output."""
        source = self.workbook()
        output = self.write('existing.xlsx', 'preserve')
        self.run_python(self.script_path, source, '--output', output, code=1)
        self.assertEqual(output.read_text(), 'preserve')
        self.run_python(self.script_path, source, '--check-only', '--output', output, code=2)
        self.run_python(self.script_path, source, '--timeout', 'nan', '--check-only', code=2)

    def test_typed_formula_errors_and_blank_caches_have_distinct_status(self):
        """CLI inspection accepts cached empty strings and flags only typed errors."""
        namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
        for with_error in (False, True):
            with self.subTest(with_error=with_error):
                book = Workbook()
                sheet = book.active
                sheet['A1'] = '=IF(TRUE,"",1)'
                sheet['A2'] = '=1/0' if with_error else '=2*3'
                sheet['A3'] = 'A note mentioning #DIV/0!'
                sheet['A4'] = '#REF!'
                sheet['A4'].data_type = 's'
                source = self.work / 'uncached.xlsx'
                book.save(source)
                book.close()
                cached = self.work / 'cached.xlsx'
                with zipfile.ZipFile(source) as original, zipfile.ZipFile(cached, 'w') as output:
                    for item in original.infolist():
                        data = original.read(item.filename)
                        if item.filename == 'xl/worksheets/sheet1.xml':
                            root = ElementTree.fromstring(data)
                            blank = root.find(f".//{namespace}c[@r='A1']")
                            blank.set('t', 'str')
                            formula = root.find(f".//{namespace}c[@r='A2']")
                            formula.set('t', 'e' if with_error else 'n')
                            formula.find(f'{namespace}v').text = '#DIV/0!' if with_error else '6'
                            data = ElementTree.tostring(root)
                        output.writestr(item, data)
                before = cached.read_bytes()
                report = self.run_python(self.script_path, cached, '--check-only',
                                         code=2 if with_error else 0).json()
                self.assertEqual(report['status'], 'errors_found' if with_error else 'success')
                self.assertEqual(report['total_formulas'], 2)
                self.assertEqual(report['missing_cache_count'], 0)
                self.assertEqual(report['errors'], [{'cell': 'Sheet!A2', 'error': '#DIV/0!'}]
                                 if with_error else [])
                self.assertEqual(report['cache_freshness'], 'not_verified')
                self.assertEqual(cached.read_bytes(), before)

    def test_macro_and_external_link_rejection_precedes_conversion(self):
        """Unsupported packages fail without publishing output or changing the input."""
        source = self.workbook()
        relationship = (b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                        b'<Relationship Id="rIdTest" Type="http://schemas.openxmlformats.org/officeDocument/'
                        b'2006/relationships/externalLink" Target="missing.xml"/>'
                        b'</Relationships>')
        for index, (member, content) in enumerate((
                ('xl/vbaProject.bin', b'synthetic macro marker'),
                ('xl/externalLinks/externalLink1.xml', b'<externalLink/>'),
                ('xl/worksheets/_rels/sheet1.xml.rels', relationship))):
            with self.subTest(member=member):
                package = self.work / f'unsupported-{index}.xlsx'
                shutil.copyfile(source, package)
                with zipfile.ZipFile(package, 'a') as archive:
                    archive.writestr(member, content)
                before = package.read_bytes()
                output = self.work / f'result-{index}.xlsx'
                report = self.run_python(self.script_path, package, '--output', output, code=1).json()
                self.assertEqual(report['status'], 'failed')
                self.assertIn('External workbook links', report['error'])
                self.assertFalse(report['recalculated'])
                self.assertFalse(output.exists())
                self.assertEqual(package.read_bytes(), before)

    @require_binary('soffice')
    def test_recalculation_preserves_formulas_and_input(self):
        """LibreOffice fills caches in a new workbook while retaining formula locations."""
        source = self.workbook()
        before = source.read_bytes()
        output = self.work / 'calculated.xlsx'
        result = self.run_python(self.script_path, source, '--output', output).json()
        self.assertEqual(result['status'], 'success')
        self.assertEqual(result['total_formulas'], 1)
        self.assertEqual(result['missing_cache_count'], 0)
        self.assertEqual(source.read_bytes(), before)
        values = load_workbook(output, data_only=True)
        formulas = load_workbook(output, data_only=False)
        try:
            self.assertEqual(values.active['A2'].value, 6)
            self.assertEqual(formulas.active['A2'].value, '=A1*3')
        finally:
            values.close()
            formulas.close()

    @require_binary('soffice')
    def test_recalculated_formula_error_has_nonzero_exit_and_readable_output(self):
        """Calculation errors are delivered for inspection with a failing CLI status."""
        source = self.work / 'division.xlsx'
        book = Workbook()
        book.active['A1'] = '=1/0'
        book.save(source)
        book.close()
        before = source.read_bytes()
        output = self.work / 'division-calculated.xlsx'
        report = self.run_python(self.script_path, source, '--output', output, code=2).json()
        self.assertEqual(report['status'], 'errors_found')
        self.assertEqual(report['errors'], [{'cell': 'Sheet!A1', 'error': '#DIV/0!'}])
        self.assertEqual(report['missing_cache_count'], 0)
        self.assertTrue(report['recalculated'])
        values = load_workbook(output, data_only=True)
        try:
            self.assertEqual(values.active['A1'].data_type, 'e')
            self.assertEqual(values.active['A1'].value, '#DIV/0!')
        finally:
            values.close()
        self.assertEqual(source.read_bytes(), before)


class GifBuilderTests(ScriptCase):
    def test_composition_deduplication_and_delivery_limits(self):
        """Composed frames retain total timing after GIF deduplication and decode correctly."""
        output = self.work / 'loop.gif'
        result = self.import_snippet('''
from core.frame_composer import create_blank_frame
from core.gif_builder import GIFBuilder
from core.validators import inspect_gif
import json, sys
builder = GIFBuilder(width=32, height=32, fps=15)
for i in range(30):
    builder.add_frame(create_blank_frame(32, 32, (255, 0, 0) if i < 15 else (0, 0, 255)))
print(json.dumps(builder.save(sys.argv[1])))
try:
    builder.save(sys.argv[1])
except FileExistsError:
    print('preserved')
else:
    raise AssertionError('overwrote existing GIF')
''', output, path_entries=[SKILLS / 'gif-creator'])
        report = json.loads(result.stdout.splitlines()[0])
        self.assertEqual(report['duration_ms'], 2000)
        self.assertEqual(report['frame_count'], 2)
        self.assertIn('preserved', result.stdout)
        result = self.run_python(SKILLS / 'gif-creator/core/validators.py', output,
                                 '--max-duration-ms', 1000, code=1).json()
        self.assertIn('duration exceeds target', result['errors'])


class MotionTests(ScriptCase):
    scripts = SKILLS / 'create-motion-videos/scripts'

    def test_copied_skill_runs_without_repository_or_sibling_resources(self):
        """A standalone copy resolves its catalog, bundled assets and project scripts."""
        installed = self.work / 'installed' / 'create-motion-videos'
        shutil.copytree(self.scripts.parent, installed,
                        ignore=shutil.ignore_patterns('__pycache__', 'node_modules'))
        scripts = installed / 'scripts'
        unrelated = self.work / 'unrelated'
        unrelated.mkdir()
        self.run_python(scripts / 'verify_bundle.py', cwd=unrelated)
        rows = self.run_python(scripts / 'styles.py', 'list', '--collection', 'all', '--json',
                               cwd=unrelated).json()
        self.assertTrue(rows)
        for collection in ('lemo', 'twoclipping', 'opuscar'):
            selected = next(row for row in rows if row['collection'] == collection)
            found = self.run_python(scripts / 'styles.py', 'show', selected['id'], '--json',
                                    cwd=unrelated).json()
            self.assertEqual(found['id'], selected['id'])
            guide = self.run_python(scripts / 'styles.py', 'show', selected['id'], cwd=unrelated)
            self.assertIn(found['title'], guide.stdout)
        project = unrelated / 'film'
        self.run_python(scripts / 'prepare_project.py', project, '--style', 'whiteboard', cwd=unrelated)
        self.assertTrue((project / 'index.html').is_file())
        self.run_python(scripts / 'audit_assets.py', project, '--strict', cwd=unrelated)

    def test_bundle_and_style_catalog_resolve(self):
        """Pinned vendored sources retain their hashes and every catalog selection resolves."""
        self.run_python(self.scripts / 'verify_bundle.py')
        rows = self.run_python(self.scripts / 'styles.py', 'list', '--collection', 'all', '--json').json()
        self.assertEqual(len({r['id'] for r in rows}), len(rows))
        for collection in ('lemo', 'twoclipping', 'opuscar'):
            row = next(r for r in rows if r['collection'] == collection)
            found = self.run_python(self.scripts / 'styles.py', 'show', row['id'], '--json').json()
            self.assertEqual(found['id'], row['id'])
        self.run_python(self.scripts / 'styles.py', 'show', 'nonexistent-style', code=2)

    def test_project_preserves_treatment_and_refuses_existing_files(self):
        """Project scaffolding carries style resources and never replaces user work."""
        project = self.work / 'project'
        treatment = self.write('project/TREATMENT.md', 'Original synthetic treatment')
        self.run_python(self.scripts / 'prepare_project.py', project, '--style', 'whiteboard')
        self.assertEqual(treatment.read_text(), 'Original synthetic treatment')
        for name in ('index.html', 'STYLE.md', 'LEMO-LICENSE', 'motion-project.json'):
            self.assertTrue((project / name).is_file())
        before = (project / 'index.html').read_bytes()
        self.run_python(self.scripts / 'prepare_project.py', project, '--style', 'whiteboard', code=2)
        self.assertEqual((project / 'index.html').read_bytes(), before)
        self.run_python(self.scripts / 'prepare_project.py', self.work / 'bad', '--style', '../whiteboard', code=2)
        self.assertFalse((self.work / 'bad').exists())

    def test_guide_only_style_has_no_fictitious_resource_paths(self):
        """Ensures documentary guides create usable starters without promising absent demos."""
        project = self.work / 'documentary'
        result = self.run_python(self.scripts / 'prepare_project.py', project,
                                 '--style', 'visual-investigation').json()
        self.assertIsNone(result['resources'])
        config = json.loads((project / 'motion-project.json').read_text())
        self.assertIsNone(config['style_resources'])
        self.assertIn('guidance only', config['instructions'])
        self.assertTrue((project / 'STYLE.md').is_file())
        self.assertTrue((project / 'index.html').is_file())
        self.assertFalse((project / 'style').exists())
        self.run_python(self.scripts / 'audit_assets.py', project, '--strict')

    def test_asset_audit_distinguishes_existing_and_missing_files(self):
        """Strict audits report a missing local asset and succeed once it is supplied."""
        self.write('film/index.html', '<img src="picture.svg">')
        project = self.work / 'film'
        self.run_python(self.scripts / 'audit_assets.py', project, '--strict', code=1)
        self.write('film/picture.svg', '<svg xmlns="http://www.w3.org/2000/svg"/>')
        self.run_python(self.scripts / 'audit_assets.py', project, '--strict')

    @require_binary('node')
    def test_render_validation_runs_before_loading_the_browser(self):
        """Invalid render requests cannot create or replace the requested video."""
        self.write('film/index.html', '<script>window.DUR=1;window.seek=()=>{};</script>')
        film = self.work / 'film'
        output = self.work / 'video.mp4'
        for args in (('--fps', '0'), ('--size', '33x32'), ('--samples', '0')):
            with self.subTest(args=args):
                self.run_node(self.scripts / 'render.mjs', film, '--out', output, *args, code=1)
                self.assertFalse(output.exists())
        output.write_bytes(b'preserve')
        self.run_node(self.scripts / 'render.mjs', film, '--out', output, code=1)
        self.assertEqual(output.read_bytes(), b'preserve')


class PromptLibraryTests(ScriptCase):
    """`create-motion-videos/scripts/prompt_library.py` and the two attributed prompt libraries."""

    scripts = SKILLS / 'create-motion-videos/scripts'
    library_files = SKILLS / 'create-motion-videos/assets/prompt-library'

    def test_both_libraries_keep_their_records_and_provenance(self):
        """Every record keeps the fields the reader prints, and the pinned source travel with it."""
        for name in ('opus55-videos', 'mg-styles-15'):
            with self.subTest(library=name):
                data = json.loads((self.library_files / f'{name}.json').read_text())
                self.assertEqual(data['library'], name)
                self.assertTrue(data['items'])
                self.assertEqual(data['source']['license'], 'MIT')
                self.assertRegex(data['source']['commit'], r'^[0-9a-f]{40}$')
                ids = [item['id'] for item in data['items']]
                self.assertEqual(len(ids), len(set(ids)), 'duplicate ids')
                for item in data['items']:
                    self.assertTrue(item['prompt'].strip())
                    if name == 'opus55-videos':
                        self.assertIsInstance(item['prompt_complete'], bool)
        self.assertEqual(len(json.loads((self.library_files / 'opus55-videos.json').read_text())['items']), 703)
        self.assertEqual(len(json.loads((self.library_files / 'mg-styles-15.json').read_text())['items']), 15)

    def test_retained_licence_notices_match_their_upstream_text(self):
        """The MIT notice for each imported collection travels inside the skill."""
        for name, holder in (('LICENSE-opus55-videos.txt', 'yihui-dev'), ('LICENSE-mg-styles-15.txt', 'Vincentwei1021')):
            text = (self.library_files / name).read_text()
            with self.subTest(notice=name):
                self.assertIn('MIT License', text)
                self.assertIn(holder, text)
                self.assertIn('without restriction', text)

    def test_list_filters_by_library_technique_and_text(self):
        """Ensures library, technique and free-text filters select exactly the matching records."""
        shader_rows = self.run_python(self.scripts / 'prompt_library.py', 'list', '--library', 'opus55-videos',
                                      '--technique', 'shader', '--json').json()
        self.assertTrue(shader_rows)
        self.assertTrue(all('shader' in row['techniques'] for row in shader_rows))
        matched = self.run_python(self.scripts / 'prompt_library.py', 'list', '--query', 'paperclip', '--json').json()
        self.assertTrue(matched)
        self.assertTrue(all('paperclip' in json.dumps(row).casefold() for row in matched))

    def test_show_prints_the_prompt_and_unknown_ids_fail(self):
        """Ensures a full record keeps its prompt, creator and post link, and that an unknown id exits 2."""
        shown = self.run_python(self.scripts / 'prompt_library.py', 'show', '09-bauhaus')
        self.assertIn('Bauhaus', shown.stdout)
        self.assertIn('# Style', shown.stdout)
        self.assertIn('```text', shown.stdout)
        attributed = self.run_python(self.scripts / 'prompt_library.py', 'show', 'ror-fly-880547')
        self.assertIn('@Ror_Fly', attributed.stdout)
        self.assertIn('https://x.com/Ror_Fly/status/', attributed.stdout)
        self.run_python(self.scripts / 'prompt_library.py', 'show', 'nonexistent-prompt', code=2)

    def test_unmatched_filters_and_unknown_libraries_are_reported(self):
        """Ensures filters that match nothing and unknown libraries fail instead of printing an empty list."""
        self.run_python(self.scripts / 'prompt_library.py', 'list', '--query', 'zzz-no-such-technique', code=2)
        self.run_python(self.scripts / 'prompt_library.py', 'list', '--library', 'nonesuch', code=2)

    def test_libraries_reports_counts_and_sources(self):
        """Ensures the libraries command reports entry counts and the pinned source of each collection."""
        report = self.run_python(self.scripts / 'prompt_library.py', 'libraries')
        self.assertIn('opus55-videos: 703 entries', report.stdout)
        self.assertIn('mg-styles-15: 15 entries', report.stdout)
        self.assertIn('@ 92aee2f586ac (MIT)', report.stdout)

    def test_a_copied_skill_resolves_the_libraries_from_any_directory(self):
        """The reader resolves assets next to itself, not from the repository cwd."""
        installed = self.work / 'installed' / 'create-motion-videos'
        shutil.copytree(self.scripts.parent, installed, ignore=shutil.ignore_patterns('__pycache__', 'node_modules'))
        unrelated = self.work / 'unrelated'
        unrelated.mkdir()
        rows = self.run_python(installed / 'scripts' / 'prompt_library.py', 'list', '--library', 'mg-styles-15',
                               '--json', cwd=unrelated).json()
        self.assertEqual(len(rows), 15)


@unittest.skipUnless(os.name == 'posix', 'server helper requires POSIX')
class ServerTests(ScriptCase):
    def test_invalid_server_arguments_never_launch_the_command(self):
        """Mismatched ports and nonfinite timeouts fail before launching user commands."""
        script = SKILLS / 'webapp-testing/scripts/with_server.py'
        marker = self.work / 'unexpected'
        command = [sys.executable, '-c', f'from pathlib import Path; Path({str(marker)!r}).touch()']
        for args in (['--server', 'exit 0', '--port', '0'],
                     ['--server', 'exit 0', '--port', '12345', '--timeout', 'nan']):
            with self.subTest(args=args):
                self.run_python(script, *args, '--', *command, code=2)
                self.assertFalse(marker.exists())


if __name__ == '__main__':
    unittest.main()
