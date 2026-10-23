"""Contract tests for the consolidated ``documents`` toolkit.

Adapted from jiahuei/agent-skills PR #8; synthetic fixtures exercise CLI
results, report shapes, OOXML preservation, and failure paths.
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import textwrap
import unittest
import zipfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from docx import Document
from docx.shared import Inches, Pt
from lxml import etree
from openpyxl import Workbook

from skill_scripts import (
    ROOT,
    ScriptCase,
    binary,
    require_module,
    require_modules,
)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"


def w(tag):
    return f"{{{W}}}{tag}"


def document_xml(path):
    with zipfile.ZipFile(path) as zf:
        return etree.fromstring(zf.read("word/document.xml"))


def part_names(path):
    with zipfile.ZipFile(path) as zf:
        return set(zf.namelist())


def paragraph_texts(path):
    return [p.text for p in Document(str(path)).paragraphs]


class DocxFixtureMixin:
    """Builders for the small .docx inputs the suites share."""

    def make_doc(self, name="base.docx", *, zoom_fix=True):
        """A document that trips the auditors on purpose.

        Heading 1 followed by Heading 3 (a level jump), a paragraph holding an
        email address and a phone number, one directly formatted run, a
        headerless table and a ``{{TAG}}`` placeholder.
        """
        doc = Document()
        doc.add_heading("Title Heading", level=1)
        doc.add_paragraph(
            "Intro paragraph with ada@example.com and +1 (555) 123-4567 inside."
        )
        doc.add_heading("Sub Heading", level=3)
        run = doc.add_paragraph().add_run("Direct formatted run")
        run.bold = True
        run.font.size = Pt(13)
        run.font.name = "Arial"
        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"
        for row, values in enumerate((("H1", "H2"), ("a", "b"))):
            for col, value in enumerate(values):
                table.cell(row, col).text = value
        doc.add_paragraph("{{TAG}} placeholder paragraph")
        if zoom_fix:
            self._fix_zoom(doc)
        path = self.work / name
        doc.save(str(path))
        return path

    @staticmethod
    def _fix_zoom(doc):
        """python-docx writes ``<w:zoom>`` without the schema-required percent."""
        zoom = doc.settings.element.find(w("zoom"))
        if zoom is not None:
            zoom.set(w("percent"), "100")

    def make_image_doc(self, name="img.docx"):
        from PIL import Image
        import io

        buf = io.BytesIO()
        Image.new("RGB", (40, 30), (10, 120, 200)).save(buf, format="PNG")
        buf.seek(0)
        doc = Document()
        doc.add_paragraph("pic doc")
        doc.add_picture(buf, width=Inches(2))
        self._fix_zoom(doc)
        path = self.work / name
        doc.save(str(path))
        return path

    def make_toc_doc(self, name="tocin.docx"):
        doc = Document()
        doc.add_paragraph("[[TOC]]")
        doc.add_heading("Alpha", level=1)
        doc.add_paragraph("body")
        doc.add_heading("Beta", level=2)
        self._fix_zoom(doc)
        path = self.work / name
        doc.save(str(path))
        return path


class DocumentScriptContract(DocxFixtureMixin):
    """Behaviour the consolidated Word toolkit must provide.

    Subclasses set :attr:`PKG`; every test resolves the script inside that
    package containing the consolidated toolkit.
    """

    PKG = None

    def setUp(self):
        super().setUp()
        self.scripts = ROOT / "skills" / self.PKG / "scripts"

    def tool(self, name):
        path = self.scripts / name
        self.assertTrue(path.is_file(), f"{self.PKG} is missing {name}")
        return path

    def cli(self, name, *args, code=0, **kwargs):
        return self.run_python(self.tool(name), *args, code=code, **kwargs)

    # -- read-only audits ------------------------------------------------
    def test_heading_audit_reports_style_counts_and_level_jumps(self):
        """Ensures heading audit reports style counts and level jumps."""
        doc = self.make_doc()
        out = self.cli("heading_audit.py", doc).stdout
        self.assertAllIn(
            ["HEADING STYLE COUNTS", "- Heading 1: 1", "- Heading 3: 1",
             "POTENTIAL HEADING LEVEL JUMPS", "Heading 1 → Heading 3", "Sub Heading"],
            out,
        )

    def test_heading_audit_max_findings_caps_the_jump_list(self):
        """Ensures heading audit max findings caps the jump list."""
        doc = Document()
        for index in range(4):
            doc.add_heading(f"H{index}", level=1 if index % 2 == 0 else 3)
        self._fix_zoom(doc)
        jumps = self.work / "jumps.docx"
        doc.save(str(jumps))
        full = self.cli("heading_audit.py", jumps).stdout
        capped = self.cli("heading_audit.py", jumps, "--max_findings", 1).stdout
        self.assertEqual(full.count("Heading 1 → Heading 3"), 2)
        self.assertEqual(capped.count("Heading 1 → Heading 3"), 1)

    def test_style_lint_summarises_and_writes_the_same_json(self):
        """Ensures style lint summarises and writes the same json."""
        doc = self.make_doc()
        report = self.work / "lint.json"
        text = self.cli("style_lint.py", doc, "--json", report).stdout
        payload = self.read_json(report)
        self.assertEqual(payload["direct_run_formatting_runs"], 1)
        self.assertEqual(payload["direct_paragraph_formatting_paragraphs"], 0)
        self.assertEqual(payload["fonts_by_char_count"], {"Arial": 20})
        suspect = payload["heading_like_paragraphs_not_heading_style"][0]
        self.assertEqual(suspect["text"], "Direct formatted run")
        self.assertEqual(suspect["style"], "Normal")
        self.assertIn("direct run-formatting runs: 1", text)
        self.assertIn("Arial(20)", text)

    def test_style_lint_reads_tables_and_headers(self):
        """Ensures style lint reads tables and headers."""
        doc = Document()
        doc.sections[0].header.paragraphs[0].add_run("Head").bold = True
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).paragraphs[0].add_run("Cell").italic = True
        self._fix_zoom(doc)
        path = self.work / "hf.docx"
        doc.save(str(path))
        report = self.work / "hf.json"
        self.cli("style_lint.py", path, "--json", report)
        self.assertEqual(self.read_json(report)["direct_run_formatting_runs"], 2)

    def test_section_audit_reports_geometry_and_linkage(self):
        """Ensures section audit reports geometry and linkage."""
        out = self.cli("section_audit.py", self.make_doc()).stdout
        self.assertAllIn(
            ["SECTIONS: 1", "orientation=PORTRAIT", "page_size(in): 8.50 x 11.00",
             "margins(in): L=1.25 R=1.25 T=1.00 B=1.00",
             "header_linked_to_previous=True", "footer_linked_to_previous=True",
             "different_first_page=False", "REMINDER"],
            out,
        )

    def test_images_audit_separates_inline_from_floating(self):
        """Ensures images audit separates inline from floating."""
        empty = self.cli("images_audit.py", self.make_doc()).stdout
        self.assertIn("No inline/anchored drawings found", empty)

        out = self.cli("images_audit.py", self.make_image_doc()).stdout
        self.assertAllIn(
            ["IMAGE KINDS", "- inline: 1", "ROWS (part | kind | size | target)",
             "word/document.xml | inline | 2.00 x 1.50 | media/image1.png"],
            out,
        )
        self.assertNotIn("WARNING", out)

    def test_images_audit_max_rows_truncates_with_a_remainder_note(self):
        """Ensures images audit max rows truncates with a remainder note."""
        from PIL import Image
        import io

        doc = Document()
        for _ in range(3):
            buf = io.BytesIO()
            Image.new("RGB", (8, 8), (1, 2, 3)).save(buf, format="PNG")
            buf.seek(0)
            doc.add_picture(buf, width=Inches(1))
        self._fix_zoom(doc)
        path = self.work / "many.docx"
        doc.save(str(path))
        out = self.cli("images_audit.py", path, "--max_rows", 2).stdout
        self.assertEqual(out.count("| inline |"), 2)
        self.assertIn("... (1 more)", out)

    def test_footnotes_report_flags_missing_parts(self):
        """Ensures footnotes report flags missing parts."""
        out = self.cli("footnotes_report.py", self.make_doc()).stdout
        self.assertAllIn(
            ["[document.xml] footnoteReferences=0 ids=[]",
             "[document.xml] endnoteReferences=0 ids=[]",
             "[word/footnotes.xml] MISSING", "[word/endnotes.xml] MISSING"],
            out,
        )

    def test_fields_report_is_quiet_without_fields(self):
        """Ensures fields report is quiet without fields."""
        out = self.cli("fields_report.py", self.make_doc()).stdout
        self.assertIn("No fields found in document.xml/headers/footers.", out)

    def test_a11y_audit_json_matches_the_summary_line(self):
        """Ensures a11y audit json matches the summary line."""
        doc = self.make_doc()
        report = self.work / "a11y.json"
        summary = self.cli("a11y_audit.py", doc, "--out_json", report).stdout
        payload = self.read_json(report)
        self.assertEqual(payload["counts"], {"high": 0, "medium": 2, "low": 0})
        self.assertEqual(
            [f["kind"] for f in payload["findings"]],
            ["heading_skip", "table_no_header_row"],
        )
        self.assertEqual(payload["findings"][0]["context"]["text"], "Sub Heading")
        self.assertIn("high=0 medium=2 low=0", summary)
        # With --out_json the report body stays out of stdout.
        self.assertNotIn("heading_skip", summary)

    def test_a11y_audit_exits_nonzero_only_for_high_severity(self):
        """Documented exit codes: 0 = no high findings, 1 = high findings present."""
        medium_only = self.work / "medium.json"
        self.cli("a11y_audit.py", self.make_doc(), "--out_json", medium_only, code=0)
        self.assertEqual(self.read_json(medium_only)["counts"]["high"], 0)

        high = self.work / "high.json"
        self.cli("a11y_audit.py", self.make_image_doc(), "--out_json", high, code=1)
        self.assertEqual(self.read_json(high)["counts"]["high"], 1)

    def test_a11y_audit_reports_missing_alt_text_as_high(self):
        """Ensures a11y audit reports missing alt text as high."""
        report = self.work / "img.json"
        self.cli("a11y_audit.py", self.make_image_doc(), "--out_json", report, code=1)
        payload = self.read_json(report)
        self.assertEqual(payload["counts"]["high"], 1)
        self.assertEqual(payload["findings"][0]["kind"], "image_missing_alt")
        self.assertEqual(payload["findings"][0]["context"]["name"], "Picture 1")

    def test_a11y_fixes_are_opt_in_and_idempotent(self):
        """Ensures a11y fixes are opt in and idempotent."""
        doc = self.make_doc()
        untouched = self.work / "untouched.docx"
        self.cli("a11y_audit.py", doc, "--out", untouched)
        self.assertEqual(
            document_xml(doc).findall(f".//{w('tblHeader')}"), [],
            "a report-only run must not mark header rows",
        )

        fixed = self.work / "fixed.docx"
        out = self.cli("a11y_audit.py", doc, "--fix_table_headers", "first_row",
                       "--fix_image_alt", "from_filename", "--out", fixed).stdout
        self.assertIn("'table_headers_set': 1", out)
        self.assertEqual(len(document_xml(fixed).findall(f".//{w('tblHeader')}")), 1)

        report = self.work / "after.json"
        self.cli("a11y_audit.py", fixed, "--out_json", report)
        kinds = [f["kind"] for f in self.read_json(report)["findings"]]
        self.assertNotIn("table_no_header_row", kinds)
        self.assertIn("heading_skip", kinds)

    def test_a11y_alt_text_fix_derives_from_the_filename(self):
        """Ensures a11y alt text fix derives from the filename."""
        fixed = self.work / "altfixed.docx"
        self.cli("a11y_audit.py", self.make_image_doc(), "--fix_image_alt",
                 "from_filename", "--out", fixed)
        descr = document_xml(fixed).find(f".//{{{WP}}}docPr")
        self.assertIsNotNone(descr)
        self.assertEqual(descr.get("descr"), "Image: image1.png")
        # The fix is reported in the counter line and clears the high finding.
        report = self.work / "altfixed.json"
        self.cli("a11y_audit.py", fixed, "--out_json", report)
        self.assertEqual(self.read_json(report)["counts"]["high"], 0)

    def test_table_to_csv_exports_the_requested_table(self):
        """Ensures table to csv exports the requested table."""
        out_path = self.work / "t.csv"
        result = self.cli("docx_table_to_csv.py", self.make_doc(), "--out", out_path)
        self.assertIn("Exported table 0", result.stdout)
        self.assertEqual(out_path.read_text().splitlines(), ["H1,H2", "a,b"])

    def test_table_to_csv_rejects_an_out_of_range_index(self):
        """Ensures table to csv rejects an out of range index."""
        result = self.cli("docx_table_to_csv.py", self.make_doc(),
                          "--table_index", 5, "--out", self.work / "no.csv",
                          code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse((self.work / "no.csv").exists())

    def test_xlsx_to_docx_table_copies_values_and_bolds_headers(self):
        """Ensures xlsx to docx table copies values and bolds headers."""
        sheet = self.work / "data.xlsx"
        book = Workbook()
        active = book.active
        active.title = "Data"
        active.append(["Name", "Qty"])
        active.append(["Apple", 3])
        active.append(["Pear", 7])
        book.create_sheet("Other").append(["ignored", "rows"])
        book.save(str(sheet))

        out = self.work / "table.docx"
        result = self.cli("xlsx_to_docx_table.py", sheet, "--out", out,
                          "--header_rows", 1, "--title", "Fruit")
        self.assertIn("rows=3, cols=2, sheet='Data'", result.stdout)

        produced = Document(str(out))
        self.assertEqual([p.text for p in produced.paragraphs if p.text], ["Fruit"])
        rows = [[c.text for c in row.cells] for row in produced.tables[0].rows]
        self.assertEqual(rows, [["Name", "Qty"], ["Apple", "3"], ["Pear", "7"]])
        self.assertTrue(all(run.bold for run in produced.tables[0].rows[0].cells[0]
                            .paragraphs[0].runs))
        self.assertFalse(all(run.bold for run in produced.tables[0].rows[1].cells[0]
                             .paragraphs[0].runs))

    def test_xlsx_to_docx_table_honours_sheet_selection(self):
        """Ensures xlsx to docx table honours sheet selection."""
        sheet = self.work / "two.xlsx"
        book = Workbook()
        book.active.title = "First"
        book.active.append(["a"])
        book.create_sheet("Second").append(["b"])
        book.save(str(sheet))
        out = self.work / "second.docx"
        self.cli("xlsx_to_docx_table.py", sheet, "--out", out, "--sheet", "Second")
        self.assertEqual(Document(str(out)).tables[0].rows[0].cells[0].text, "b")
        self.cli("xlsx_to_docx_table.py", sheet, "--out", self.work / "missing.docx",
                 "--sheet", "Nope", code=None).returncode
        self.assertFalse((self.work / "missing.docx").exists())

    # -- privacy / redaction ---------------------------------------------
    def test_privacy_scrub_strips_metadata_but_leaves_body_text_alone(self):
        """Ensures privacy scrub strips metadata but leaves body text alone."""
        doc = Document()
        doc.add_paragraph("Contact ada@example.com")
        doc.core_properties.author = "Someone"
        doc.core_properties.last_modified_by = "Someone"
        self._fix_zoom(doc)
        source = self.work / "meta.docx"
        doc.save(str(source))

        out = self.work / "scrubbed.docx"
        result = self.cli("privacy_scrub.py", source, "--out", out)
        self.assertIn("[OK] wrote", result.stdout)
        stats = re.search(r"\{.*\}", result.stdout).group(0)
        self.assertGreaterEqual(json.loads(stats.replace("'", '"'))["core_props_scrubbed"], 1)

        self.assertEqual(paragraph_texts(out), ["Contact ada@example.com"],
                         "privacy_scrub is metadata-only; redaction is a separate tool")
        self.assertNotIn(b"Someone", zipfile.ZipFile(out).read("docProps/core.xml"))
        self.assertGreater(self._rsid_count(source), self._rsid_count(out))
        self.assertEqual(self._rsid_count(out), 0, "every rsid attribute must be gone")

    @staticmethod
    def _rsid_count(path):
        with zipfile.ZipFile(path) as zf:
            return sum(len(re.findall(rb"w:rsid\w*=", zf.read(name)))
                       for name in zf.namelist() if name.endswith(".xml"))

    def test_redact_masks_matches_and_reports_counts(self):
        """Ensures redact masks matches and reports counts."""
        source = self.make_doc()
        out = self.work / "redacted.docx"
        report = self.cli("redact_docx.py", source, out, "--emails", "--phones").stdout
        payload = json.loads(report)
        self.assertEqual(payload["parts_processed"], 1)
        self.assertEqual(payload["paragraphs_touched"], 1)
        self.assertEqual(payload["matches_redacted"], 2)

        text = paragraph_texts(out)[1]
        self.assertNotIn("ada@example.com", text)
        self.assertNotIn("123-4567", text)
        self.assertIn("█" * len("ada@example.com"), text)
        self.assertIn("Intro paragraph with", text, "surrounding text is preserved")

    def test_redact_supports_custom_patterns_mask_and_replacement(self):
        """Ensures redact supports custom patterns mask and replacement."""
        source = self.make_doc()
        masked = self.work / "masked.docx"
        self.cli("redact_docx.py", source, masked, "--pattern", r"Intro",
                 "--mask_char", "X")
        self.assertIn("XXXXX", paragraph_texts(masked)[1])

        replaced = self.work / "replaced.docx"
        self.cli("redact_docx.py", source, replaced, "--pattern", r"ada@example\.com",
                 "--replacement", "[REDACTED]", "--no_preserve_length")
        self.assertIn("[REDACTED]", paragraph_texts(replaced)[1])

    def test_redact_can_reach_comments_and_leaves_them_alone_by_default(self):
        """Ensures redact can reach comments and leaves them alone by default."""
        commented = self.work / "commented.docx"
        self.cli("comments_add.py", self.make_doc(), "--out", commented,
                 "--add", "Intro=note with ada@example.com")
        default = self.work / "no_comments.docx"
        self.cli("redact_docx.py", commented, default, "--emails")
        self.assertIn("ada@example.com",
                      zipfile.ZipFile(default).read("word/comments.xml").decode())

        included = self.work / "with_comments.docx"
        self.cli("redact_docx.py", commented, included, "--emails", "--include_comments")
        self.assertNotIn("ada@example.com",
                         zipfile.ZipFile(included).read("word/comments.xml").decode())

    # -- comments ---------------------------------------------------------
    def test_comment_lifecycle_add_extract_patch_strip(self):
        """Ensures comment lifecycle add extract patch strip."""
        source = self.make_doc()
        added = self.work / "added.docx"
        result = self.cli("comments_add.py", source, "--out", added, "--author", "QA",
                          "--add", "Intro=a comment on intro")
        self.assertIn("added_comments=1, unmatched_patterns=0", result.stdout)

        listing = self.work / "comments.json"
        self.cli("comments_extract.py", added, "--out", listing)
        payload = self.read_json(listing)
        self.assertEqual(payload["comment_count"], 1)
        comment = payload["comments"][0]
        self.assertEqual(comment["author"], "QA")
        self.assertEqual(comment["text"], "a comment on intro")
        self.assertFalse(comment["resolved"])
        self.assertRegex(comment["date"], r"^\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}Z$")
        anchors = comment["anchors"][0]
        self.assertEqual(anchors["part"], "word/document.xml")
        self.assertEqual([p["where"] for p in anchors["paragraphs"]], ["start", "end"])

        patched = self.work / "patched.docx"
        self.write_json("patch.json", {"ops": [
            {"id": 0, "append": "Follow-up"},
            {"id": 0, "resolved": True},
        ]})
        result = self.cli("comments_apply_patch.py", added, self.work / "patch.json",
                          "--out", patched)
        self.assertIn("comments_touched=1 ops_applied=2", result.stdout)
        self.cli("comments_extract.py", patched, "--out", self.work / "patched.json")
        updated = self.read_json(self.work / "patched.json")["comments"][0]
        self.assertEqual(updated["text"], "a comment on intro\nFollow-up")
        self.assertTrue(updated["resolved"])

        stripped = self.work / "stripped.docx"
        result = self.cli("comments_strip.py", patched, "--out", stripped)
        self.assertAllIn(["markup_removed", "comments_part_removed': 1",
                          "rels_updated': 1", "content_types_updated': 1"], result.stdout)
        self.assertNotIn("word/comments.xml", part_names(stripped))
        self.cli("comments_extract.py", stripped, "--out", self.work / "empty.json")
        self.assertEqual(self.read_json(self.work / "empty.json")["comment_count"], 0)

    def test_comment_patch_replace_clears_previous_paragraphs(self):
        """Ensures comment patch replace clears previous paragraphs."""
        added = self.work / "added.docx"
        self.cli("comments_add.py", self.make_doc(), "--out", added,
                 "--add", "Intro=original")
        self.write_json("replace.json", {"ops": [{"id": 0, "replace": "One\nTwo"}]})
        out = self.work / "replaced.docx"
        self.cli("comments_apply_patch.py", added, self.work / "replace.json", "--out", out)
        self.cli("comments_extract.py", out, "--out", self.work / "replaced.json")
        self.assertEqual(self.read_json(self.work / "replaced.json")["comments"][0]["text"],
                         "One\nTwo")

    def test_comment_patch_unknown_id_warns_and_applies_nothing(self):
        """Ensures comment patch unknown id warns and applies nothing."""
        added = self.work / "added.docx"
        self.cli("comments_add.py", self.make_doc(), "--out", added, "--add", "Intro=x")
        self.write_json("unknown.json", {"ops": [{"id": 99, "append": "nope"}]})
        out = self.work / "unknown.docx"
        result = self.cli("comments_apply_patch.py", added, self.work / "unknown.json",
                          "--out", out)
        self.assertIn("comment id 99 not found", result.stderr + result.stdout)
        self.assertIn("comments_touched=0 ops_applied=0", result.stdout)

    def test_comments_add_require_all_turns_a_miss_into_a_failure(self):
        """Ensures comments add require all turns a miss into a failure."""
        out = self.work / "required.docx"
        result = self.cli("comments_add.py", self.make_doc(), "--out", out,
                          "--add", "NoSuchText=c", "--require_all", code=1)
        self.assertIn("no paragraph matched contains='NoSuchText'",
                      result.stdout + result.stderr)
        self.assertIn("1 patterns were not matched", result.stdout + result.stderr)
        listing = self.work / "required.json"
        self.cli("comments_extract.py", out, "--out", listing)
        self.assertEqual(self.read_json(listing)["comment_count"], 0)

        # Without --require_all the same miss is a warning and a clean exit.
        lenient = self.work / "lenient.docx"
        result = self.cli("comments_add.py", self.make_doc(), "--out", lenient,
                          "--add", "NoSuchText=c", "--add", "Intro=kept")
        self.assertIn("added_comments=1, unmatched_patterns=1", result.stdout)

    def test_comments_add_ignore_case_matches_upper_case_input(self):
        """Ensures comments add ignore case matches upper case input."""
        strict = self.work / "strict.docx"
        result = self.cli("comments_add.py", self.make_doc(), "--out", strict,
                          "--add", "INTRO PARAGRAPH=c", "--require_all", code=1)
        self.assertIn("no paragraph matched", result.stdout + result.stderr)

        lenient = self.work / "ci.docx"
        result = self.cli("comments_add.py", self.make_doc(), "--out", lenient,
                          "--add", "INTRO PARAGRAPH=c", "--ignore_case")
        self.assertIn("added_comments=1", result.stdout)
        listing = self.work / "ci.json"
        self.cli("comments_extract.py", lenient, "--out", listing)
        self.assertEqual(self.read_json(listing)["comment_count"], 1)

    def test_comments_extract_on_a_document_without_comments(self):
        """Ensures comments extract on a document without comments."""
        source = self.make_doc()
        out = self.work / "none.json"
        self.cli("comments_extract.py", source, "--out", out)
        self.assertEqual(self.read_json(out),
                         {"file": str(source), "comment_count": 0, "comments": []})

    # -- content controls --------------------------------------------------
    def test_content_control_list_wrap_and_fill_round_trip(self):
        """Ensures content control list wrap and fill round trip."""
        source = self.make_doc()
        self.assertIn("no SDTs found", self.cli("content_controls.py", source, "list").stdout)

        wrapped = self.work / "wrapped.docx"
        result = self.cli("content_controls.py", source, "wrap_placeholders",
                          "--output", wrapped)
        self.assertIn("wrapped placeholders in 1 run(s)", result.stdout)
        listing = self.cli("content_controls.py", wrapped, "list").stdout
        self.assertAllIn(["word/document.xml", "tag=TAG", "alias=TAG", "text={{TAG}}"],
                         listing)
        as_json = self.cli("content_controls.py", wrapped, "list", "--json").stdout
        self.assertEqual(json.loads(as_json),
                         [{"part": "word/document.xml", "tag": "TAG",
                           "alias": "TAG", "text": "{{TAG}}"}])

        filled = self.work / "filled.docx"
        result = self.cli("content_controls.py", wrapped, "fill", "--set", "TAG=HELLO",
                          "--output", filled)
        self.assertIn("filled 1 SDT(s)", result.stdout)
        # python-docx does not descend into w:sdt, so read the part directly.
        with zipfile.ZipFile(filled) as zf:
            body = zf.read("word/document.xml").decode()
        self.assertIn('<w:sdtContent><w:r><w:t>HELLO</w:t></w:r></w:sdtContent>', body)
        self.assertNotIn("{{TAG}}", body)
        self.assertIn(" placeholder paragraph", body)

    def test_content_controls_are_visible_in_the_sdt_xml(self):
        """Ensures content controls are visible in the sdt xml."""
        wrapped = self.work / "wrapped.docx"
        self.cli("content_controls.py", self.make_doc(), "wrap_placeholders",
                 "--output", wrapped)
        sdt = document_xml(wrapped).find(f".//{w('sdt')}")
        self.assertIsNotNone(sdt)
        self.assertEqual(sdt.find(f".//{w('tag')}").get(w("val")), "TAG")

    # -- fields, TOC, captions --------------------------------------------
    def test_insert_toc_requires_an_exact_placeholder_paragraph(self):
        """Ensures insert toc requires an exact placeholder paragraph."""
        missing = self.cli("insert_toc.py", self.make_doc(), "--out", self.work / "x.docx",
                           code=2)
        self.assertIn("Could not find a paragraph whose text == placeholder '[[TOC]]'",
                      missing.stdout + missing.stderr)

        source = self.make_toc_doc()
        out = self.work / "toc.docx"
        self.cli("insert_toc.py", source, "--out", out, "--levels", "1-2")
        fields = self.cli("fields_report.py", out).stdout
        self.assertIn("- TOC: 1", fields)
        self.assertAllIn(["FIELD COUNTS", "EXAMPLES", "PER-PART DETAILS", "== word/document.xml ==",
                          "REMINDER"], fields)
        instr = document_xml(out).find(f".//{w('instrText')}").text
        self.assertIn(r'TOC \o "1-2"', instr)

    def test_captions_bookmarks_and_ref_fields_then_materialize(self):
        """Ensures captions bookmarks and ref fields then materialize."""
        captioned = self.work / "caps.docx"
        self.cli("captions_and_crossrefs.py", self.make_doc(), captioned, "--tables",
                 "--bookmarks", "--caption_text", "Results")
        self.assertIn("- SEQ: 1", self.cli("fields_report.py", captioned).stdout)
        self.assertIn("SEQ Table", self.cli("fields_report.py", captioned).stdout)
        self.assertIsNotNone(document_xml(captioned).find(f".//{w('bookmarkStart')}"))

        marked = self.work / "marked.docx"
        doc = Document(str(captioned))
        doc.add_paragraph("See [[REF:tbl1]] now.")
        self._fix_zoom(doc)
        doc.save(str(marked))

        refs = self.work / "refs.docx"
        result = self.cli("insert_ref_fields.py", marked, refs)
        self.assertIn("markers_replaced=1", result.stdout)
        self.assertNotIn("[[REF:tbl1]]", " ".join(paragraph_texts(refs)))

        materialized = self.work / "materialized.docx"
        result = self.cli("fields_materialize.py", refs, "--out", materialized)
        # Both the caption's SEQ field and the new REF field get a cached result.
        self.assertIn("materialized 2 field(s)", result.stdout)
        text = " ".join(paragraph_texts(materialized))
        self.assertIn("See", text)
        self.assertIn("Table 1", text)
        # Materializing writes the cached result but keeps the live field, so
        # Word can still update it; the markup therefore survives the rewrite.
        with zipfile.ZipFile(materialized) as zf:
            body = zf.read("word/document.xml").decode()
        self.assertIn("instrText", body)

    def test_insert_ref_fields_reports_markers_split_across_runs(self):
        """Ensures insert ref fields reports markers split across runs."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("See [[REF:")
        para.add_run("tbl1]] now.")
        self._fix_zoom(doc)
        source = self.work / "split.docx"
        doc.save(str(source))
        out = self.work / "split_out.docx"
        result = self.cli("insert_ref_fields.py", source, out)
        self.assertIn("markers_replaced=0", result.stdout)
        self.assertIn("[[REF:tbl1]]", " ".join(paragraph_texts(out)))

    def test_flatten_ref_fields_removes_the_field_markup(self):
        """Ensures flatten ref fields removes the field markup."""
        captioned = self.work / "caps.docx"
        self.cli("captions_and_crossrefs.py", self.make_doc(), captioned, "--tables",
                 "--bookmarks")

        # A caption alone holds only a SEQ field, which flatten_ref_fields ignores.
        untouched = self.work / "untouched.docx"
        self.assertIn("fields_flattened=0",
                      self.cli("flatten_ref_fields.py", captioned, "--out",
                               untouched).stdout)

        marked = self.work / "marked.docx"
        doc = Document(str(captioned))
        doc.add_paragraph("See [[REF:tbl1]] now.")
        self._fix_zoom(doc)
        doc.save(str(marked))
        refs = self.work / "refs.docx"
        self.cli("insert_ref_fields.py", marked, refs)

        flat = self.work / "flat.docx"
        result = self.cli("flatten_ref_fields.py", refs, "--out", flat)
        self.assertIn("fields_flattened=1", result.stdout)
        with zipfile.ZipFile(flat) as zf:
            body = zf.read("word/document.xml").decode()
        self.assertNotIn("REF tbl1", body)
        self.assertIn("See", " ".join(paragraph_texts(flat)))

    def test_fields_materialize_only_flag_limits_what_is_rewritten(self):
        """Ensures fields materialize only flag limits what is rewritten."""
        captioned = self.work / "caps.docx"
        self.cli("captions_and_crossrefs.py", self.make_doc(), captioned, "--tables")
        out = self.work / "only_ref.docx"
        result = self.cli("fields_materialize.py", captioned, "--out", out, "--only", "REF")
        self.assertIn("materialized 0 field(s)", result.stdout)
        self.assertIn("SEQ Table", self.cli("fields_report.py", out).stdout)

    def test_insert_note_creates_footnotes_and_endnotes(self):
        """Ensures insert note creates footnotes and endnotes."""
        footnote = self.work / "fn.docx"
        self.cli("insert_note.py", self.make_doc(), "--kind", "footnote",
                 "--text", "A note", "--out", footnote)
        report = self.cli("footnotes_report.py", footnote).stdout
        self.assertAllIn(
            ["[document.xml] footnoteReferences=1 ids=[1]",
             "[word/footnotes.xml] defined_footnotes=1 ids=[1]",
             "id=1: A note", "[word/endnotes.xml] MISSING"],
            report,
        )

        endnote = self.work / "en.docx"
        self.cli("insert_note.py", self.make_doc(), "--kind", "endnote",
                 "--text", "End matter", "--out", endnote)
        report = self.cli("footnotes_report.py", endnote).stdout
        self.assertAllIn(["[document.xml] endnoteReferences=1",
                          "[word/endnotes.xml] defined_endnotes=1", "End matter"], report)

    def test_internal_navigation_adds_toc_and_back_links(self):
        """Ensures internal navigation adds toc and back links."""
        out = self.work / "nav.docx"
        self.cli("internal_nav.py", self.make_toc_doc(), "--out", out, "--toc_title",
                 "Contents", "--levels", "1", "2")
        texts = paragraph_texts(out)
        self.assertIn("Contents", texts)
        self.assertIn("Alpha", texts)
        xml = etree.tostring(document_xml(out)).decode()
        self.assertIn('w:bookmarkStart', xml)
        self.assertIn('w:name="TOC"', xml)
        # Navigation is real internal hyperlinks, not HYPERLINK field codes.
        self.assertIn('w:hyperlink w:anchor="sec001"', xml)
        self.assertIn('w:hyperlink w:anchor="TOC"', xml)
        self.assertIn("Alpha Back to TOC", texts)
        self.assertIn("Beta Back to TOC", texts)
        self.assertIn("Top | Bottom | TOC", texts)

        quiet = self.work / "quiet.docx"
        self.cli("internal_nav.py", self.make_toc_doc(), "--out", quiet,
                 "--no_quicklinks", "--no_back_to_toc", "--no_top_bottom")
        quiet_texts = paragraph_texts(quiet)
        self.assertNotIn("Top | Bottom | TOC", quiet_texts)
        self.assertNotIn("Alpha Back to TOC", " ".join(quiet_texts))
        # The opt-out flags suppress the extras but still emit the static TOC.
        self.assertIn("Alpha", quiet_texts)

    # -- tracked changes ----------------------------------------------------
    def test_tracked_change_report_accept_and_reject(self):
        """Ensures tracked change report accept and reject."""
        source = self.make_doc()
        self.assertIn("[report] ins=0 del=0 moveTo=0 moveFrom=0",
                      self.cli("accept_tracked_changes.py", source, "--mode", "report").stdout)

        tracked = self.work / "tracked.docx"
        result = self.cli("add_tracked_replacements.py", source, "--out", tracked,
                          "--author", "Bot", "--replace", "Intro=Introduction")
        self.assertIn("replacements=1", result.stdout)
        self.assertIn("[report] ins=1 del=1",
                      self.cli("accept_tracked_changes.py", tracked, "--mode",
                               "report").stdout)

        accepted = self.work / "accepted.docx"
        result = self.cli("accept_tracked_changes.py", tracked, "--mode", "accept",
                          "--out", accepted)
        self.assertIn("[after]  ins=0 del=0", result.stdout)
        self.assertTrue(paragraph_texts(accepted)[1].startswith("Introduction paragraph"))

        rejected = self.work / "rejected.docx"
        self.cli("accept_tracked_changes.py", tracked, "--mode", "reject",
                 "--out", rejected)
        self.assertFalse(paragraph_texts(rejected)[1].startswith("Introduction"))
        self.assertIn("paragraph with ada@example.com", paragraph_texts(rejected)[1])

    def test_tracked_change_accept_requires_an_output_path(self):
        """Ensures tracked change accept requires an output path."""
        tracked = self.work / "tracked.docx"
        self.cli("add_tracked_replacements.py", self.make_doc(), "--out", tracked,
                 "--replace", "Intro=Introduction")
        result = self.cli("accept_tracked_changes.py", tracked, "--mode", "accept", code=1)
        self.assertIn("--out is required for accept/reject", result.stdout + result.stderr)

    def test_tracked_change_accept_keeps_tracking_when_asked(self):
        """Ensures tracked change accept keeps tracking when asked."""
        tracked = self.work / "tracked.docx"
        self.cli("add_tracked_replacements.py", self.make_doc(), "--out", tracked,
                 "--replace", "Intro=Introduction")
        kept = self.work / "kept.docx"
        self.cli("accept_tracked_changes.py", tracked, "--mode", "accept",
                 "--keep_tracking_on", "--out", kept)
        with zipfile.ZipFile(kept) as zf:
            self.assertIn(b"trackRevisions", zf.read("word/settings.xml"))
        dropped = self.work / "dropped.docx"
        self.cli("accept_tracked_changes.py", tracked, "--mode", "accept", "--out", dropped)
        with zipfile.ZipFile(dropped) as zf:
            self.assertNotIn(b"trackRevisions", zf.read("word/settings.xml"))

    def test_make_fixtures_generates_the_documented_inputs(self):
        """Ensures make fixtures generates the documented inputs."""
        outdir = self.work / "fixtures"
        self.cli("make_fixtures.py", "--outdir", outdir, "--only", "all")
        self.assertEqual(
            sorted(p.name for p in outdir.iterdir()),
            ["tracked_changes_fixture.docx", "watermark_fixture.docx"],
        )
        self.assertIn("w:ins", etree.tostring(
            document_xml(outdir / "tracked_changes_fixture.docx")).decode())

        only = self.work / "only"
        self.cli("make_fixtures.py", "--outdir", only, "--only", "watermark")
        self.assertEqual([p.name for p in only.iterdir()], ["watermark_fixture.docx"])

    # -- structure edits ------------------------------------------------------
    def test_merge_append_concatenates_bodies_in_order(self):
        """Ensures merge append concatenates bodies in order."""
        base = self.work / "base.docx"
        first = Document()
        first.add_paragraph("Base para")
        self._fix_zoom(first)
        first.save(str(base))
        extra = self.work / "extra.docx"
        second = Document()
        second.add_paragraph("Appended para")
        self._fix_zoom(second)
        second.save(str(extra))

        merged = self.work / "merged.docx"
        result = self.cli("merge_docx_append.py", base, extra, "--out", merged)
        self.assertIn("body_children_appended': 1", result.stdout)
        self.assertEqual(paragraph_texts(merged), ["Base para", "Appended para"])
        # Inputs are never rewritten in place.
        self.assertEqual(paragraph_texts(base), ["Base para"])

    def test_style_normalize_reports_the_overrides_it_clears(self):
        """Ensures style normalize reports the overrides it clears."""
        out = self.work / "normalized.docx"
        result = self.cli("style_normalize.py", self.make_doc(), out,
                          "--enforce_heading_spacing")
        self.assertAllIn(["run_overrides_cleared=1", "para_overrides_cleared=0",
                          "heading_spacing_updates=2"], result.stdout)
        produced = Document(str(out))
        bold_run = produced.paragraphs[3].runs[0]
        self.assertIsNone(bold_run.bold)
        self.assertEqual(produced.paragraphs[0].paragraph_format.space_after, Pt(6))

    def test_style_normalize_positional_and_flag_outputs_agree(self):
        """Ensures style normalize positional and flag outputs agree."""
        positional = self.work / "pos.docx"
        flagged = self.work / "flag.docx"
        source = self.make_doc()
        self.cli("style_normalize.py", source, positional)
        self.cli("style_normalize.py", source, "--out", flagged,
                 "--clear_paragraph_format")
        self.assertTrue(positional.is_file() and flagged.is_file())

    def test_apply_template_styles_copies_style_definitions(self):
        """Ensures apply template styles copies style definitions."""
        template = self.work / "template.docx"
        doc = Document()
        doc.styles["Normal"].font.name = "Georgia"
        self._fix_zoom(doc)
        doc.save(str(template))

        target = self.work / "target.docx"
        plain = Document()
        plain.add_paragraph("Body copy")
        self._fix_zoom(plain)
        plain.save(str(target))

        out = self.work / "styled.docx"
        self.cli("apply_template_styles.py", "--template", template, "--target", target,
                 "--out", out)
        with zipfile.ZipFile(out) as zf:
            self.assertIn(b"Georgia", zf.read("word/styles.xml"))

    def test_set_protection_modes_and_aliases(self):
        """Ensures set protection modes and aliases."""
        for mode, expected in (("comments", b'w:edit="comments"'),
                               ("read_only", b'w:edit="readOnly"'),
                               ("tracked-changes", b'w:edit="trackedChanges"'),
                               ("forms", b'w:edit="forms"')):
            out = self.work / f"prot-{mode}.docx"
            self.cli("set_protection.py", self.make_doc(), "--mode", mode, "--out", out)
            with zipfile.ZipFile(out) as zf:
                self.assertIn(expected, zf.read("word/settings.xml"))

        off = self.work / "prot-off.docx"
        self.cli("set_protection.py", self.make_doc(), "--mode", "off", "--out", off)
        with zipfile.ZipFile(off) as zf:
            self.assertNotIn(b"documentProtection", zf.read("word/settings.xml"))

    def test_set_protection_rejects_unknown_modes(self):
        """Ensures set protection rejects unknown modes."""
        out = self.work / "bad.docx"
        result = self.cli("set_protection.py", self.make_doc(), "--mode", "bogus",
                          "--out", out, code=1)
        self.assertIn("invalid --mode='bogus'", result.stdout + result.stderr)
        self.assertFalse(out.exists())

    def test_watermark_round_trip(self):
        """Ensures watermark round trip."""
        marked = self.work / "wm.docx"
        result = self.cli("watermark_add.py", self.make_doc(), "--out", marked,
                          "--text", "DRAFT")
        self.assertIn("patched word/header1.xml", result.stdout)

        report = self.cli("watermark_audit_remove.py", marked, "--mode", "report").stdout
        self.assertAllIn(["[word/header1.xml]", "vml_textpath: DRAFT",
                          "attr_watermark:", "[summary] watermark_like_hits=2"], report)

        cleaned = self.work / "clean.docx"
        self.cli("watermark_audit_remove.py", marked, "--mode", "remove",
                 "--contains", "DRAFT", "--out", cleaned)
        after = self.cli("watermark_audit_remove.py", cleaned, "--mode", "report").stdout
        self.assertIn("watermark_like_hits=0", after)
        self.assertNotIn(b"DRAFT", zipfile.ZipFile(cleaned).read("word/header1.xml"))

    def test_watermark_removal_matches_the_stamp_text_case_insensitively(self):
        """Ensures watermark removal matches the stamp text case insensitively."""
        marked = self.work / "wm.docx"
        self.cli("watermark_add.py", self.make_doc(), "--out", marked, "--text", "DRAFT")

        lowered = self.work / "lowered.docx"
        self.cli("watermark_audit_remove.py", marked, "--mode", "remove",
                 "--contains", "draft", "--out", lowered)
        self.assertIn("watermark_like_hits=0",
                      self.cli("watermark_audit_remove.py", lowered,
                               "--mode", "report").stdout)

        # A filter that matches no stamp text leaves the package untouched.
        noop = self.work / "noop.docx"
        self.cli("watermark_audit_remove.py", marked, "--mode", "remove",
                 "--contains", "absent", "--out", noop)
        self.assertIn("watermark_like_hits=2",
                      self.cli("watermark_audit_remove.py", noop,
                               "--mode", "report").stdout)

    def test_watermark_removal_needs_both_a_filter_and_an_output(self):
        """Ensures watermark removal needs both a filter and an output."""
        marked = self.work / "wm.docx"
        self.cli("watermark_add.py", self.make_doc(), "--out", marked, "--text", "DRAFT")
        result = self.cli("watermark_audit_remove.py", marked, "--mode", "remove",
                          code=2)
        self.assertIn("--contains and --out are required when --mode remove",
                      result.stdout + result.stderr)

    def test_docx_ooxml_patch_applies_the_requested_operations(self):
        """Ensures docx ooxml patch applies the requested operations."""
        out = self.work / "patched.docx"
        result = self.cli("docx_ooxml_patch.py", self.make_doc(), "--out", out,
                          "--enable-track", "--add-page-numbers", "--author", "QA")
        self.assertAllIn(["[OK] Patched ->", "add-page-numbers", "enable-track"],
                         result.stdout)
        with zipfile.ZipFile(out) as zf:
            settings = zf.read("word/settings.xml")
            footer = zf.read("word/footer1.xml")
        self.assertIn(b"trackRevisions", settings)
        self.assertIn(b"PAGE", footer)

    def test_docx_ooxml_patch_defaults_to_in_place(self):
        """Ensures docx ooxml patch defaults to in place."""
        source = self.make_doc()
        before = source.read_bytes()
        self.cli("docx_ooxml_patch.py", source, "--enable-track")
        self.assertNotEqual(source.read_bytes(), before)
        with zipfile.ZipFile(source) as zf:
            self.assertIn(b"trackRevisions", zf.read("word/settings.xml"))

    def test_render_and_diff_reports_visual_and_structural_differences(self):
        """``--render_py`` lets the diff run without LibreOffice installed."""
        stub = self.write(
            "stub_render.py",
            textwrap.dedent(
                """
                import argparse, hashlib
                from pathlib import Path
                from PIL import Image, ImageDraw

                parser = argparse.ArgumentParser()
                parser.add_argument("input_path")
                parser.add_argument("--output_dir")
                parser.add_argument("--width", type=int)
                parser.add_argument("--height", type=int)
                parser.add_argument("--dpi", type=int)
                parser.add_argument("--emit_pdf", action="store_true")
                parser.add_argument("--verbose", action="store_true")
                args = parser.parse_args()

                out = Path(args.output_dir)
                out.mkdir(parents=True, exist_ok=True)
                seed = int(hashlib.sha256(
                    Path(args.input_path).read_bytes()).hexdigest()[:8], 16)
                for index in range(2):
                    image = Image.new("RGB", (300, 400), (255, 255, 255))
                    ImageDraw.Draw(image).rectangle(
                        [20 + index * 10 + seed % 7, 30, 200, 60],
                        fill=(seed % 255, index * 40, 90))
                    # render_docx.py names pages page-<N>.png, 1-based.
                    image.save(out / f"page-{index + 1}.png")
                """
            ),
        )
        left = self.work / "left.docx"
        first = Document()
        first.add_paragraph("Shared opening line")
        first.add_paragraph("Only in the left document")
        self._fix_zoom(first)
        first.save(str(left))
        right = self.work / "right.docx"
        second = Document()
        second.add_paragraph("Shared opening line")
        second.add_paragraph("Only in the right document")
        self._fix_zoom(second)
        second.save(str(right))

        outdir = self.work / "diff"
        result = self.cli("render_and_diff.py", left, right, "--outdir", outdir,
                          "--render_py", stub)
        self.assertIn("[OK] wrote", result.stdout)
        self.assertIn("[summary] changed_pages=[1, 2]", result.stdout)

        summary = self.read_json(outdir / "summary.json")
        self.assertEqual(summary["a_docx"], str(left))
        self.assertEqual(summary["b_docx"], str(right))
        self.assertEqual(summary["pages_a"], 2)
        self.assertEqual(summary["pages_b"], 2)
        self.assertEqual(summary["changed_pages"], [1, 2])
        self.assertEqual(summary["num_changed_pages"], 2)

        for folder in ("a_render", "b_render"):
            self.assertEqual(
                sorted(p.name for p in (outdir / folder).iterdir()),
                ["page-1.png", "page-2.png"], folder)
        # Changed pages get a visual diff next to the renders.
        self.assertEqual(
            sorted(p.name for p in (outdir / "diff_pages").iterdir()),
            ["diff-page-1.png", "diff-page-2.png"])

        # The structural diff is a unified text diff of the two bodies.
        self.assertIn("Shared opening line", (outdir / "text_a.txt").read_text())
        self.assertIn("Only in the left document", (outdir / "text_a.txt").read_text())
        self.assertIn("Only in the right document", (outdir / "text_b.txt").read_text())
        diff = (outdir / "text_diff.txt").read_text()
        self.assertIn("-Only in the left document", diff)
        self.assertIn("+Only in the right document", diff)

    def test_render_and_diff_reports_no_changes_for_identical_documents(self):
        """Ensures render and diff reports no changes for identical documents."""
        stub = self.write("noop_render.py", textwrap.dedent(
            """
            import argparse
            from pathlib import Path
            from PIL import Image

            parser = argparse.ArgumentParser()
            parser.add_argument("input_path")
            parser.add_argument("--output_dir")
            parser.add_argument("--dpi", type=int)
            parser.add_argument("--width", type=int)
            parser.add_argument("--height", type=int)
            parser.add_argument("--emit_pdf", action="store_true")
            parser.add_argument("--verbose", action="store_true")
            args = parser.parse_args()
            out = Path(args.output_dir)
            out.mkdir(parents=True, exist_ok=True)
            Image.new("RGB", (100, 100), (255, 255, 255)).save(out / "page-1.png")
            """
        ))
        source = self.make_doc()
        outdir = self.work / "same"
        result = self.cli("render_and_diff.py", source, source, "--outdir", outdir,
                          "--render_py", stub)
        self.assertIn("[summary] changed_pages=[]", result.stdout)
        self.assertEqual(self.read_json(outdir / "summary.json")["changed_pages"], [])
        self.assertEqual((outdir / "text_diff.txt").read_text().strip(), "")
        self.assertFalse((outdir / "diff_pages").exists(),
                         "no page changed, so no diff image is written")

    def test_docx_ooxml_patch_exports_the_shared_helper_surface(self):
        """Six sibling scripts import these names; keep them available."""
        self.import_snippet(
            "import docx_ooxml_patch as m\n"
            "for name in ('unzip_docx','zip_docx','iso_now','enable_track_revisions',"
            "'tracked_replace_ins_id','add_comment_to_first_indented_paragraph'):\n"
            "    assert callable(getattr(m, name)), name\n"
            "print('surface-ok')\n",
            path_entries=[self.scripts],
        ).stdout.index("surface-ok")




class DocumentsPackageTests(DocumentScriptContract, ScriptCase):
    PKG = "documents"




class RevisionHelperTests(DocxFixtureMixin, ScriptCase):
    """Revision and comment helpers consolidated into ``skills/documents``."""

    def setUp(self):
        super().setUp()
        self.scripts = ROOT / "skills/documents/scripts"

    def tool(self, name):
        return self.scripts / name

    def test_merge_runs_collapses_identical_adjacent_runs_only(self):
        """Ensures merge runs collapses identical adjacent runs only."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("one ")
        para.add_run("two ")          # same formatting -> merges
        bold = para.add_run("three")
        bold.bold = True              # different formatting -> stays separate
        self._fix_zoom(doc)
        source = self.work / "runs.docx"
        doc.save(str(source))

        out = self.work / "merged.docx"
        result = self.run_python(self.tool("merge_runs.py"), source, "-o", out)
        merged = re.search(r"Merged (\d+) runs", result.stdout)
        self.assertIsNotNone(merged, result.detail)
        self.assertEqual(int(merged.group(1)), 1)
        self.assertEqual(Document(str(out)).paragraphs[0].text, "one two three")
        self.assertEqual(len(Document(str(out)).paragraphs[0].runs), 2)
        self.assertTrue(Document(str(out)).paragraphs[0].runs[1].bold)

    def test_merge_runs_refuses_a_packed_input_without_an_output(self):
        """Ensures merge runs refuses a packed input without an output."""
        source = self.make_doc()
        result = self.run_python(self.tool("merge_runs.py"), source, code=2)
        self.assertIn("--output", result.stdout + result.stderr)

    def test_merge_runs_accepts_an_unpacked_directory_in_place(self):
        """Ensures merge runs accepts an unpacked directory in place."""
        doc = Document()
        para = doc.add_paragraph()
        para.add_run("split ")
        para.add_run("runs")
        self._fix_zoom(doc)
        source = self.work / "runs.docx"
        doc.save(str(source))

        unpacked = self.work / "unpacked"
        with zipfile.ZipFile(source) as zf:
            zf.extractall(unpacked)
        result = self.run_python(self.tool("merge_runs.py"), unpacked)
        # Directory input is edited in place, so there is no output file to name.
        self.assertEqual(result.stdout.strip(), "Merged 1 runs")
        self.assertTrue((unpacked / "word/document.xml").is_file())
        self.assertNotIn(b">split <", (unpacked / "word/document.xml").read_bytes())

    def test_comment_helper_escapes_by_default_and_supports_raw(self):
        """Ensures comment helper escapes by default and supports raw."""
        source = self.make_doc()
        escaped = self.work / "escaped.docx"
        result = self.run_python(self.tool("comment.py"), source, "Hello <b>world</b>",
                                 "-o", escaped, "--author", "QA")
        self.assertAllIn(["Added comment id=0", "commentRangeStart", "commentReference"],
                         result.stdout)
        with zipfile.ZipFile(escaped) as zf:
            body = zf.read("word/comments.xml").decode()
        self.assertIn("&lt;b&gt;world&lt;/b&gt;", body)

        raw = self.work / "raw.docx"
        self.run_python(self.tool("comment.py"), source, "Plain <b>bold</b>", "--raw",
                        "-o", raw)
        with zipfile.ZipFile(raw) as zf:
            self.assertIn("<b>bold</b>", zf.read("word/comments.xml").decode())

    def test_comment_helper_auto_assigns_the_next_id_and_replies(self):
        """Ensures comment helper auto assigns the next id and replies."""
        first = self.work / "c1.docx"
        self.run_python(self.tool("comment.py"), self.make_doc(), "first", "-o", first)
        second = self.work / "c2.docx"
        result = self.run_python(self.tool("comment.py"), first, "reply", "--parent", "0",
                                 "-o", second)
        self.assertIn("Added reply id=1", result.stdout)
        self.assertIn("Nest the reply's markers inside parent 0's", result.stdout)
        with zipfile.ZipFile(second) as zf:
            body = zf.read("word/comments.xml").decode()
            threading = zf.read("word/commentsExtended.xml").decode()
        self.assertIn('w:id="0"', body)
        self.assertIn('w:id="1"', body)
        # Replies are threaded through w15:paraIdParent, not the comments part.
        self.assertIn("w15:paraIdParent=", threading)
        self.assertEqual(threading.count("w15:paraIdParent="), 1)

    def test_comment_helper_refuses_a_duplicate_id(self):
        """Ensures comment helper refuses a duplicate id."""
        first = self.work / "c1.docx"
        self.run_python(self.tool("comment.py"), self.make_doc(), "first", "-o", first)
        result = self.run_python(self.tool("comment.py"), first, "again", "--id", "0",
                                 "-o", self.work / "dup.docx", code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertFalse((self.work / "dup.docx").exists())

    @unittest.skipUnless(binary("soffice") or binary("libreoffice"),
                         "accept_changes.py drives LibreOffice")
    def test_accept_changes_never_reports_success_on_timeout(self):
        """Ensures accept changes never reports success on timeout."""
        source = self.make_doc()
        result = self.run_python(self.tool("accept_changes.py"), source,
                                 self.work / "out.docx", "--timeout", "1", code=None)
        self.assertIn(result.returncode, (0, 1), result.detail)
        if result.returncode != 0:
            self.assertFalse((self.work / "out.docx").exists(),
                             "a failed conversion must not publish an output")


class DocumentAuditTests(DocxFixtureMixin, ScriptCase):
    """Table geometry and Google Docs title audits."""

    def setUp(self):
        super().setUp()
        self.scripts = ROOT / "skills/documents/scripts"

    def tool(self, name):
        return self.scripts / name

    def test_table_geometry_reports_per_row_widths_and_issues(self):
        """A geometry mismatch is a failure: the tool exits 1 after reporting it."""
        result = self.run_python(self.tool("table_geometry.py"), self.make_doc(), code=1)
        self.assertAllIn(
            ["table 1:", "tblW=auto:0", "grid_sum=8640", "grid=[4320, 4320]",
             "ISSUE: table width is missing or not DXA",
             "ISSUE: table indent is missing or not DXA",
             "ISSUE: grid column sum does not equal table width",
             "row 1: tcW=[4320, 4320] sum=8640", "ISSUES: 3"],
            result.stdout,
        )

    def test_table_geometry_is_clean_for_a_document_without_tables(self):
        """Ensures table geometry is clean for a document without tables."""
        doc = Document()
        doc.add_paragraph("no tables here")
        self._fix_zoom(doc)
        path = self.work / "plain.docx"
        doc.save(str(path))
        out = self.run_python(self.tool("table_geometry.py"), path).stdout
        self.assertEqual(
            out.strip(),
            "OK: all tables have matching tblW, tblInd, tblGrid, and tcW",
        )

    def test_table_geometry_accepts_an_explicit_dxa_table(self):
        """A DXA table whose indent matches the default cell margin passes clean."""
        source = self.make_doc()
        sized = self.work / "sized.docx"
        with zipfile.ZipFile(source) as zin, zipfile.ZipFile(
            sized, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    data = data.replace(
                        b'<w:tblW w:type="auto" w:w="0"/>',
                        b'<w:tblW w:type="dxa" w:w="8640"/>'
                        b'<w:tblInd w:type="dxa" w:w="120"/>',
                    )
                zout.writestr(info, data)
        out = self.run_python(self.tool("table_geometry.py"), sized).stdout
        self.assertIn("tblW=dxa:8640", out)
        self.assertIn("tblInd=dxa:120", out)
        self.assertIn("grid=[4320, 4320]", out)
        self.assertIn("OK: all tables have matching tblW, tblInd, tblGrid, and tcW", out)

    def test_table_geometry_indent_rule_follows_the_cell_margin(self):
        """An explicit tcMar/start becomes the expected indent, not the default."""
        source = self.make_doc()
        sized = self.work / "margined.docx"
        with zipfile.ZipFile(source) as zin, zipfile.ZipFile(
            sized, "w", zipfile.ZIP_DEFLATED
        ) as zout:
            for info in zin.infolist():
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    data = data.replace(
                        b'<w:tblW w:type="auto" w:w="0"/>',
                        b'<w:tblW w:type="dxa" w:w="8640"/>'
                        b'<w:tblInd w:type="dxa" w:w="0"/>',
                    ).replace(
                        b'<w:tcW w:type="dxa" w:w="4320"/>',
                        b'<w:tcW w:type="dxa" w:w="4320"/>'
                        b'<w:tcMar><w:start w:type="dxa" w:w="0"/></w:tcMar>',
                    )
                zout.writestr(info, data)
        out = self.run_python(self.tool("table_geometry.py"), sized).stdout
        self.assertIn("tblInd=dxa:0", out)
        self.assertIn("OK: all tables have matching tblW, tblInd, tblGrid, and tcW", out)
        self.assertNotIn("table indent should match", out)

    def test_google_docs_title_check_fails_on_border_residue(self):
        """Ensures google docs title check fails on border residue."""
        result = self.run_python(self.tool("google_docs_title_sanitize.py"),
                                 self.make_doc(), "--check", code=1)
        self.assertIn("Title paragraph style contains paragraph border residue",
                      result.stdout + result.stderr)

    def test_google_docs_title_sanitize_reports_each_fix_class(self):
        """Ensures google docs title sanitize reports each fix class."""
        out = self.work / "sanitized.docx"
        result = self.run_python(self.tool("google_docs_title_sanitize.py"),
                                 self.make_doc(), "--out", out)
        self.assertAllIn(["[OK] wrote", "style_borders_removed=1",
                          "paragraph_borders_removed=0", "title_underlines_removed=0"],
                         result.stdout)
        self.assertTrue(out.is_file())
        # Sanitizing twice is a no-op, so --check passes on the result.
        again = self.work / "again.docx"
        self.run_python(self.tool("google_docs_title_sanitize.py"), out, "--out", again)
        self.run_python(self.tool("google_docs_title_sanitize.py"), again, "--check")

    def test_google_docs_title_in_place_is_atomic_and_mutually_exclusive(self):
        """Ensures google docs title in place is atomic and mutually exclusive."""
        source = self.make_doc()
        before = source.read_bytes()
        self.run_python(self.tool("google_docs_title_sanitize.py"), source, "--in-place")
        self.assertNotEqual(source.read_bytes(), before)
        result = self.run_python(self.tool("google_docs_title_sanitize.py"), source,
                                 "--in-place", "--out", self.work / "x.docx", code=2)
        self.assertTrue(result.stdout + result.stderr)

    def test_google_docs_title_leading_paragraph_window_is_configurable(self):
        """Ensures google docs title leading paragraph window is configurable."""
        out = self.work / "wide.docx"
        result = self.run_python(
            self.tool("google_docs_title_sanitize.py"), self.make_doc(), "--out", out,
            "--leading-nonempty-paragraphs", "5")
        self.assertIn("[OK] wrote", result.stdout)


class DocumentPreservationTests(DocxFixtureMixin, ScriptCase):
    """Regressions for targeted ZIP editing and failure-safe document operations."""

    def setUp(self):
        super().setUp()
        self.scripts = ROOT / "skills/documents/scripts"

    def cli(self, name, *args, **kwargs):
        return self.run_python(self.scripts / name, *args, **kwargs)

    def rewrite_parts(self, source, destination, replacements):
        with zipfile.ZipFile(source) as original, zipfile.ZipFile(destination, "w") as output:
            output.comment = original.comment
            remaining = dict(replacements)
            for info in original.infolist():
                output.writestr(info, remaining.pop(info.filename, original.read(info)))
            for name, data in remaining.items():
                output.writestr(name, data)
        return destination

    def test_targeted_edits_preserve_unknown_zip_entries_and_metadata(self):
        """All package editors retain opaque extension payloads and ZIP metadata."""
        source = self.make_toc_doc()
        with zipfile.ZipFile(source, "a") as archive:
            archive.comment = b"synthetic archive annotation"
            info = zipfile.ZipInfo("customXml/opaque.bin", (2020, 1, 2, 3, 4, 6))
            info.external_attr = 0o640 << 16
            info.comment = b"opaque metadata"
            archive.writestr(info, b"opaque\x00payload")
        template = self.make_doc("template.docx")
        added = self.work / "commented.docx"
        self.cli("comments_add.py", source, "--add", "body=review", "--out", added)
        # comments_add predates this targeted rewrite; restore the ZIP annotation.
        with zipfile.ZipFile(added, "a") as archive:
            archive.comment = b"synthetic archive annotation"
        patch = self.write_json("change.json", {"ops": [{"id": 0, "append": "extra"}]})
        cases = [
            ("insert_toc.py", [source]),
            ("insert_note.py", [source, "--kind", "footnote", "--text", "A note"]),
            ("watermark_add.py", [source, "--text", "DRAFT"]),
            ("watermark_audit_remove.py", [source, "--mode", "remove", "--contains", "absent"]),
            ("google_docs_title_sanitize.py", [source]),
            ("apply_template_styles.py", ["--target", source, "--template", template]),
            ("merge_docx_append.py", [source, template]),
            ("a11y_audit.py", [source, "--fix_table_headers", "first_row"]),
            ("comments_apply_patch.py", [added, patch]),
        ]
        for index, (name, args) in enumerate(cases):
            with self.subTest(script=name):
                output = self.work / f"preserved-{index}.docx"
                self.cli(name, *args, "--out", output)
                before = added if name == "comments_apply_patch.py" else source
                with zipfile.ZipFile(before) as original, zipfile.ZipFile(output) as revised:
                    self.assertEqual(revised.comment, original.comment)
                    a = original.getinfo("customXml/opaque.bin")
                    b = revised.getinfo("customXml/opaque.bin")
                    self.assertEqual(revised.read(b), original.read(a))
                    self.assertEqual((b.date_time, b.comment, b.external_attr),
                                     (a.date_time, a.comment, a.external_attr))
                    self.assertEqual(revised.read("docProps/core.xml"), original.read("docProps/core.xml"))

    def rename_owned_part(self, source, output, kind, old_name, new_name, target):
        """Build a valid alternate OPC part location with an unchanged relationship ID."""
        with zipfile.ZipFile(source) as archive:
            rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
            for rel in rels:
                if rel.get("Type", "").endswith("/" + kind):
                    rel.set("Target", target)
            types = etree.fromstring(archive.read("[Content_Types].xml"))
            for override in types:
                if override.get("PartName") == "/" + old_name:
                    override.set("PartName", "/" + new_name)
            with zipfile.ZipFile(output, "w") as revised:
                for info in archive.infolist():
                    if info.filename == "word/_rels/document.xml.rels":
                        revised.writestr(info, etree.tostring(rels))
                    elif info.filename == "[Content_Types].xml":
                        revised.writestr(info, etree.tostring(types))
                    elif info.filename == old_name:
                        revised.writestr(new_name, archive.read(info))
                    else:
                        revised.writestr(info, archive.read(info))
        return output

    def test_toc_uses_existing_renamed_settings_and_preserves_anchors(self):
        """TOC edits preserve bookmark/comment ranges and a nonstandard settings target."""
        source = self.make_toc_doc()
        root = document_xml(source)
        paragraph = root.find(f"{w('body')}/{w('p')}")
        for local, attributes in (("bookmarkStart", {w("id"): "42", w("name"): "Contents"}),
                                  ("commentRangeStart", {w("id"): "17"})):
            paragraph.insert(0, etree.Element(w(local), attributes))
        for local, attributes in (("bookmarkEnd", {w("id"): "42"}),
                                  ("commentRangeEnd", {w("id"): "17"})):
            paragraph.append(etree.Element(w(local), attributes))
        ref = etree.SubElement(paragraph, w("r"))
        etree.SubElement(ref, w("commentReference"), {w("id"): "17"})
        anchored = self.rewrite_parts(source, self.work / "anchored.docx", {"word/document.xml": etree.tostring(root)})
        for number, target in enumerate(("settings2.xml", "/word/settings2.xml", "nested/../settings2.xml")):
            renamed = self.rename_owned_part(anchored, self.work / f"settings-{number}.docx", "settings",
                                            "word/settings.xml", "word/settings2.xml", target)
            self.cli("insert_toc.py", renamed, "--out", renamed)
            self.assertIsNotNone(Document(renamed).settings)
            with zipfile.ZipFile(renamed) as archive:
                self.assertNotIn("word/settings.xml", archive.namelist())
                self.assertIn(b"updateFields", archive.read("word/settings2.xml"))
                rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
                self.assertEqual(sum(r.get("Type", "").endswith("/settings") for r in rels), 1)
            paragraph = document_xml(renamed).find(f"{w('body')}/{w('p')}")
            tags = [child.tag for child in paragraph]
            self.assertEqual(tags[:2], [w("commentRangeStart"), w("bookmarkStart")])
            self.assertEqual(tags[-3:-1], [w("bookmarkEnd"), w("commentRangeEnd")])
            self.assertEqual(len(paragraph.findall(f".//{w('commentReference')}")), 1)

    def test_toc_inserts_field_update_in_schema_order(self):
        """A new updateFields setting follows CT_Settings order in the bundled XSD."""
        source = self.make_toc_doc()
        schema = etree.parse(str(self.scripts / "office/schemas/ecma376-transitional/wml.xsd"))
        namespaces = {"xs": "http://www.w3.org/2001/XMLSchema"}
        sequence = schema.find(".//xs:complexType[@name='CT_Settings']/xs:sequence", namespaces)
        order = [node.get("name") or node.get("ref", "").split(":")[-1] for node in sequence]
        for suffix, late_settings in (("early", ""), ("mixed", '<w:compat/><w:rsids/><w:decimalSymbol w:val="."/>')):
            settings = f'<w:settings xmlns:w="{W}"><w:zoom w:percent="100"/>{late_settings}</w:settings>'
            fixture = self.rewrite_parts(source, self.work / f"settings-{suffix}.docx", {
                "word/settings.xml": settings.encode()})
            self.cli("insert_toc.py", fixture, "--out", fixture)
            with zipfile.ZipFile(fixture) as archive:
                saved = etree.fromstring(archive.read("word/settings.xml"))
            indices = [order.index(etree.QName(node).localname) for node in saved]
            self.assertEqual(indices, sorted(indices))
            self.assertEqual(saved.find(w("updateFields")).get(w("val")), "true")

    def test_notes_and_styles_reuse_their_relationship_targets(self):
        """Existing singleton parts keep their paths and relationship IDs after edits."""
        source = self.make_doc()
        noted = self.work / "noted.docx"
        self.cli("insert_note.py", source, "--kind", "footnote", "--text", "first", "--out", noted)
        renamed = self.rename_owned_part(noted, self.work / "custom-notes.docx", "footnotes",
                                        "word/footnotes.xml", "word/notes/footnotes.xml", "notes/footnotes.xml")
        self.cli("insert_note.py", renamed, "--kind", "footnote", "--text", "second", "--out", renamed)
        with zipfile.ZipFile(renamed) as archive:
            self.assertNotIn("word/footnotes.xml", archive.namelist())
            notes = etree.fromstring(archive.read("word/notes/footnotes.xml"))
            self.assertEqual([n.get(w("id")) for n in notes][-2:], ["1", "2"])
            rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
            self.assertEqual(sum(r.get("Type", "").endswith("/footnotes") for r in rels), 1)
        template_doc = Document()
        template_doc.styles["Normal"].font.name = "Georgia"
        template = self.work / "brand.docx"
        template_doc.save(template)
        template = self.rename_owned_part(template, self.work / "custom-brand.docx", "styles",
                                          "word/styles.xml", "word/brand/styles.xml", "/word/brand/styles.xml")
        target = self.rename_owned_part(source, self.work / "custom-target.docx", "styles",
                                        "word/styles.xml", "word/styles2.xml", "styles2.xml")
        self.cli("apply_template_styles.py", "--template", template, "--target", target, "--out", target)
        self.assertEqual(Document(target).styles["Normal"].font.name, "Georgia")
        self.assertNotIn("word/styles.xml", part_names(target))

    def test_toc_rejects_ambiguous_and_external_settings_relationships(self):
        """Invalid singleton relationships fail before a destination is replaced."""
        source = self.make_toc_doc()
        with zipfile.ZipFile(source) as archive:
            original = archive.read("word/_rels/document.xml.rels")
        for fault in ("duplicate", "external", "missing"):
            rels = etree.fromstring(original)
            settings = next(r for r in rels if r.get("Type", "").endswith("/settings"))
            if fault == "duplicate":
                duplicate = etree.fromstring(etree.tostring(settings))
                duplicate.set("Id", "rId999")
                rels.append(duplicate)
            elif fault == "external":
                settings.set("TargetMode", "External")
                settings.set("Target", "https://example.org/settings.xml")
            else:
                settings.set("Target", "missing.xml")
            broken = self.rewrite_parts(source, self.work / f"{fault}.docx", {
                "word/_rels/document.xml.rels": etree.tostring(rels)})
            destination = self.write_bytes("existing-output.docx", b"preserve")
            self.cli("insert_toc.py", broken, "--out", destination, code=2)
            self.assertEqual(destination.read_bytes(), b"preserve")

    def test_atomic_package_save_keeps_destination_when_serialization_fails(self):
        """A ZIP write failure cannot truncate an existing destination or leak staging files."""
        source = self.make_doc()
        destination = self.write_bytes("existing.docx", b"keep existing destination")
        snippet = textwrap.dedent(f"""
            from unittest.mock import patch
            from word_package import Package
            package = Package({str(source)!r})
            with patch("word_package.zipfile.ZipFile.writestr", side_effect=OSError("simulated disk error")):
                try:
                    package.save({str(destination)!r})
                except OSError:
                    pass
                else:
                    raise AssertionError("write error swallowed")
            """)
        self.import_snippet(snippet, path_entries=[self.scripts])
        self.assertEqual(destination.read_bytes(), b"keep existing destination")
        self.assertEqual(list(self.work.glob(".docx-*.tmp")), [])

    def test_toc_missing_token_or_invalid_levels_does_not_publish_output(self):
        """TOC preconditions are checked before replacing a destination."""
        destination = self.write_bytes("existing.docx", b"unchanged")
        self.cli("insert_toc.py", self.make_doc(), "--out", destination, code=2)
        self.assertEqual(destination.read_bytes(), b"unchanged")
        self.cli("insert_toc.py", self.make_toc_doc(), "--levels", "3-1", "--out", destination, code=2)
        self.assertEqual(destination.read_bytes(), b"unchanged")

    def test_note_marker_order_and_section_boundary_survive(self):
        """Notes replace one token in place and append fallback references before sectPr."""
        doc = Document()
        doc.add_paragraph("Before [[NOTE]] after [[NOTE]] end").runs[0].bold = True
        source = self.work / "marker.docx"
        doc.save(source)
        self.cli("insert_note.py", source, "--kind", "footnote", "--text", "one", "--out", source)
        paragraph = document_xml(source).find(f"{w('body')}/{w('p')}")
        run = paragraph.find(w("r"))
        self.assertEqual([node.tag for node in run], [w("rPr"), w("t"), w("footnoteReference"), w("t")])
        self.assertEqual(run.findall(w("t"))[0].text, "Before ")
        self.assertEqual(run.findall(w("t"))[1].text, " after [[NOTE]] end")
        self.assertEqual(run.findall(w("t"))[0].get("{http://www.w3.org/XML/1998/namespace}space"), "preserve")
        self.cli("insert_note.py", source, "--kind", "footnote", "--text", "two", "--marker", "absent", "--out", source)
        body = document_xml(source).find(w("body"))
        self.assertEqual(body[-1].tag, w("sectPr"))
        self.assertEqual(body[-2].find(f".//{w('footnoteReference')}").get(w("id")), "2")
        with zipfile.ZipFile(source) as archive:
            notes = etree.fromstring(archive.read("word/footnotes.xml"))
            self.assertEqual([n.get(w("id")) for n in notes], ["-1", "0", "1", "2"])
            self.assertIn(b"footnotes", archive.read("word/_rels/document.xml.rels"))

    def test_content_wrap_preserves_breaks_formatting_and_split_text_nodes(self):
        """Tokens in adjacent text nodes wrap without discarding other run children."""
        doc = Document()
        paragraph = doc.add_paragraph()
        run = paragraph.add_run("prefix ")
        run.bold = True
        run.add_break()
        run.add_text("{{TA")
        run.add_text("G}} suffix")
        source = self.work / "tokens.docx"
        doc.save(source)
        before = source.read_bytes()
        output = self.work / "wrapped.docx"
        self.cli("content_controls.py", source, "wrap_placeholders", "--output", output)
        root = document_xml(output)
        self.assertEqual(len(root.findall(f".//{w('br')}")), 1)
        self.assertEqual(len(root.findall(f".//{w('sdt')}")), 1)
        self.assertIsNotNone(root.find(f".//{w('sdtContent')}/{w('r')}/{w('rPr')}/{w('b')}"))
        self.assertEqual("".join(root.itertext()), "prefix {{TAG}} suffix")
        self.assertEqual(source.read_bytes(), before)
        # Running wrap again must not turn the placeholder inside the control into a nested SDT.
        self.cli("content_controls.py", output, "wrap_placeholders", "--output", output)
        self.assertEqual(len(document_xml(output).findall(f".//{w('sdt')}")), 1)
        self.cli("content_controls.py", output, "fill", "--set", "TAG= value ", "--output", output)
        text = document_xml(output).find(f".//{w('sdtContent')}/{w('r')}/{w('t')}")
        self.assertEqual(text.text, " value ")
        self.assertEqual(text.get("{http://www.w3.org/XML/1998/namespace}space"), "preserve")

    def test_content_fill_rejects_bad_assignments_without_touching_output(self):
        """Invalid mapping syntax fails before a destination is replaced."""
        output = self.write_bytes("destination.docx", b"preserve")
        self.cli("content_controls.py", self.make_doc(), "fill", "--set", "=value", "--output", output, code=2)
        self.assertEqual(output.read_bytes(), b"preserve")

    def test_a11y_image_fix_resolves_header_relationships_locally(self):
        """Header image IDs are resolved from the header's own relationships part."""
        source = self.make_image_doc()
        with zipfile.ZipFile(source) as archive:
            drawing = document_xml(source).find(f".//{w('drawing')}")
            drawing_xml = etree.tostring(drawing)
        header = f'<w:hdr xmlns:w="{W}"><w:p><w:r>'.encode() + drawing_xml + b'</w:r></w:p></w:hdr>'
        embed = drawing.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}blip").get(
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
        rels = f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="{embed}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="media/header-logo.png"/></Relationships>'.encode()
        custom = self.rewrite_parts(source, self.work / "header.docx", {
            "word/header1.xml": header, "word/_rels/header1.xml.rels": rels})
        output = self.work / "header-fixed.docx"
        self.cli("a11y_audit.py", custom, "--fix_image_alt", "from_filename", "--out", output)
        with zipfile.ZipFile(output) as archive:
            properties = etree.fromstring(archive.read("word/header1.xml")).find(f".//{{{WP}}}docPr")
            self.assertEqual(properties.get("descr"), "Image: header-logo.png")

    def test_table_geometry_is_usable_through_python_docx_and_validates_before_mutation(self):
        """Library edits retain python-docx element types and refuse merged cells unchanged."""
        snippet = textwrap.dedent("""
            from docx import Document
            from lxml import etree
            from table_geometry import apply_table_geometry, column_widths_from_weights
            document = Document()
            table = document.add_table(rows=2, cols=2)
            widths = column_widths_from_weights([1, 2], 9000)
            apply_table_geometry(table, widths)
            assert [column.width.twips for column in table.columns] == [3000, 6000]
            assert [cell.width.twips for cell in table.rows[0].cells] == [3000, 6000]
            table.cell(0, 0).merge(table.cell(0, 1))
            before = etree.tostring(table._tbl)
            try:
                apply_table_geometry(table, [4000, 5000])
            except ValueError:
                pass
            else:
                raise AssertionError("merged table accepted")
            assert etree.tostring(table._tbl) == before
            """)
        self.import_snippet(snippet, path_entries=[self.scripts])

    def test_template_style_part_dependencies_are_refused_before_output(self):
        """A style pack with asset relationships cannot silently create dangling links."""
        source = self.make_doc()
        rels = b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="font" Target="fonts/example.odttf"/></Relationships>'
        template = self.rewrite_parts(source, self.work / "font-template.docx", {"word/_rels/fontTable.xml.rels": rels})
        output = self.write_bytes("destination.docx", b"preserve")
        result = self.cli("apply_template_styles.py", "--template", template, "--target", source, "--out", output, code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("full package merger", result.stderr)
        self.assertEqual(output.read_bytes(), b"preserve")

    def test_merge_refuses_drawings_and_preserves_base_final_section(self):
        """Unsafe drawing append fails without replacing output; plain append preserves geometry."""
        source = self.make_doc()
        output = self.write_bytes("combined.docx", b"preserve")
        result = self.cli("merge_docx_append.py", source, self.make_image_doc(), "--out", output, code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(output.read_bytes(), b"preserve")
        original = etree.tostring(document_xml(source).find(f"{w('body')}/{w('sectPr')}"))
        self.cli("merge_docx_append.py", source, self.make_toc_doc(), "--out", output)
        body = document_xml(output).find(w("body"))
        self.assertEqual(body[-1].tag, w("sectPr"))
        self.assertEqual(etree.tostring(body[-1]), original)

    def test_comment_patch_missing_part_and_malformed_ops_preserve_destination(self):
        """Invalid patches cannot replace a caller's existing output."""
        source = self.make_doc()
        output = self.write_bytes("patched.docx", b"preserve")
        for name, payload in (("no-comments.json", {"ops": []}), ("bad-ops.json", {"ops": [None]})):
            patch = self.write_json(name, payload)
            result = self.cli("comments_apply_patch.py", source, patch, "--out", output, code=None)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(output.read_bytes(), b"preserve")

    def test_comments_extract_sorts_mixed_ids_and_preserves_paragraph_boundaries(self):
        """Numeric and nonnumeric comment IDs remain exportable in a deterministic order."""
        source = self.make_doc()
        comments = etree.Element(w("comments"), nsmap={"w": W})
        for identifier in ("label", "10", "2"):
            comment = etree.SubElement(comments, w("comment"), {w("id"): identifier})
            for text in ("First paragraph", "Second paragraph"):
                paragraph = etree.SubElement(comment, w("p"))
                run = etree.SubElement(paragraph, w("r"))
                etree.SubElement(run, w("t")).text = text
        source = self.rewrite_parts(source, self.work / "mixed-comments.docx", {
            "word/comments.xml": etree.tostring(comments)})
        output = self.work / "comments.json"
        self.cli("comments_extract.py", source, "--out", output)
        entries = self.read_json(output)["comments"]
        self.assertEqual([entry["id"] for entry in entries], ["2", "10", "label"])
        self.assertEqual(entries[0]["text"], "First paragraph\nSecond paragraph")

    def test_title_sanitizer_preserves_borders_after_the_title_window(self):
        """Only selected leading paragraphs lose borders; later body styling survives."""
        doc = Document()
        for index in range(5):
            paragraph = doc.add_paragraph(f"Line {index}")
            properties = paragraph._p.get_or_add_pPr()
            border = etree.SubElement(properties, w("pBdr"))
            etree.SubElement(border, w("bottom"), {w("val"): "single"})
            paragraph.add_run(" underlined").underline = True
        source = self.work / "borders.docx"
        doc.save(source)
        output = self.work / "clean-borders.docx"
        self.cli("google_docs_title_sanitize.py", source, "--out", output,
                 "--leading-nonempty-paragraphs", "2")
        paragraphs = document_xml(output).find(w("body")).findall(w("p"))
        self.assertEqual([p.find(f"{w('pPr')}/{w('pBdr')}") is not None for p in paragraphs],
                         [False, False, True, True, True])
        self.assertEqual(len(document_xml(output).findall(f".//{w('u')}")), 5)

    def test_watermark_removal_preserves_unmatched_pictures_in_the_same_header(self):
        """The text filter must not remove another stamp or neighboring header text."""
        source = self.make_doc()
        output = self.work / "stamps.docx"
        self.cli("watermark_add.py", source, "--text", "DRAFT", "--out", output)
        self.cli("watermark_add.py", output, "--text", "KEEP BRAND", "--out", output)
        self.cli("watermark_audit_remove.py", output, "--mode", "remove", "--contains", "draft", "--out", output)
        with zipfile.ZipFile(output) as archive:
            header = etree.fromstring(archive.read("word/header1.xml"))
        paths = header.findall(".//{urn:schemas-microsoft-com:vml}textpath")
        self.assertEqual([path.get("string") for path in paths], ["KEEP BRAND"])
        self.assertEqual(len(header.findall(f".//{w('pict')}")), 1)

    def test_render_diff_rejects_empty_renderer_output(self):
        """A renderer exiting successfully without pages must not generate a success report."""
        renderer = self.write("empty.py", "# A deliberately empty renderer\n")
        output = self.work / "empty-result"
        source = self.make_doc()
        result = self.cli("render_and_diff.py", source, source, "--outdir", output, "--render_py", renderer, code=None)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("no page PNGs", result.stderr)
        self.assertFalse((output / "summary.json").exists())

    def test_render_diff_rerun_removes_stale_generated_pages(self):
        """Fresh staging prevents old pages or diff images contaminating a later comparison."""
        renderer = self.write("one_page.py", textwrap.dedent("""
            import sys
            from pathlib import Path
            from PIL import Image
            destination = Path(sys.argv[sys.argv.index("--output_dir") + 1])
            Image.new("RGB", (10, 10), "white").save(destination / "page-1.png")
            """))
        output = self.work / "comparison"
        (output / "a_render").mkdir(parents=True)
        (output / "a_render/page-99.png").write_bytes(b"stale")
        (output / "diff_pages").mkdir()
        (output / "diff_pages/diff-page-99.png").write_bytes(b"stale")
        (output / "user-note.txt").write_text("keep")
        source = self.make_doc()
        self.cli("render_and_diff.py", source, source, "--outdir", output, "--render_py", renderer)
        self.assertEqual(self.read_json(output / "summary.json")["pages_a"], 1)
        self.assertFalse((output / "a_render/page-99.png").exists())
        self.assertFalse((output / "diff_pages").exists())
        self.assertEqual((output / "user-note.txt").read_text(), "keep")


class OfficeValidatorLibraryTests(DocxFixtureMixin, ScriptCase):
    """``scripts/office/`` — the validator library behind ``validate.py``.

    ``office/validators/base.py`` holds the shared OPC/schema checks,
    ``office/validators/docx.py`` the Word-specific ones and
    ``office/validators/redlining.py`` the body-text comparison against
    ``--original``.  They are imported directly here so the library surface the
    CLI is built on is pinned independently of the command line.
    """

    def setUp(self):
        super().setUp()
        self.scripts = ROOT / "skills/documents/scripts"

    def library(self, body):
        return self.import_snippet(textwrap.dedent(body),
                                   path_entries=[self.scripts]).stdout

    def unpack(self, docx, name="unpacked"):
        target = self.work / name
        with zipfile.ZipFile(docx) as zf:
            zf.extractall(target)
        return target

    def test_the_shared_base_validator_exposes_the_opc_checks(self):
        """Ensures the shared base validator exposes the opc checks."""
        docx = self.unpack(self.make_doc())
        out = self.library(
            f"""
            import json
            from office.helpers import XML_SPACE
            from office.validators.base import PackageValidator
            validator = PackageValidator({str(docx)!r})
            names = [name for name in dir(validator) if name.startswith("check_")]
            print(json.dumps(sorted(names)))
            print(XML_SPACE == " " + chr(9) + chr(13) + chr(10))
            print(validator.check_well_formed())
            try:
                validator.validate()
            except NotImplementedError as exc:
                print("abstract:", exc)
            """
        )
        lines = out.splitlines()
        checks = json.loads(lines[0])
        for expected in ("check_well_formed", "check_namespaces", "check_id_uniqueness",
                         "check_part_references", "check_relationship_id_refs",
                         "check_content_types", "check_xsd_delta"):
            self.assertIn(expected, checks)
        self.assertEqual(lines[1], "True")
        self.assertEqual(lines[2], "True", "check_well_formed() reports success as a bool")
        self.assertIn("abstract: family validators implement validate()",
                      lines[3], "the base class stays abstract")

    def test_the_docx_validator_adds_the_word_specific_checks(self):
        """Ensures the docx validator adds the word specific checks."""
        docx = self.unpack(self.make_doc())
        out = self.library(
            f"""
            import json
            from office.validators.base import PackageValidator
            from office.validators.docx import DOCXSchemaValidator
            print(issubclass(DOCXSchemaValidator, PackageValidator))
            validator = DOCXSchemaValidator({str(docx)!r})
            base = {{n for n in dir(PackageValidator) if n.startswith("check_")}}
            extra = sorted({{n for n in dir(validator) if n.startswith("check_")}} - base)
            print(json.dumps(extra))
            validator.report_paragraph_delta()
            """
        )
        lines = out.splitlines()
        self.assertEqual(lines[0], "True")
        extra = json.loads(lines[1])
        for expected in ("check_whitespace_preservation", "check_deleted_markup",
                         "check_inserted_markup", "check_revision_id_ranges",
                         "check_comment_anchoring"):
            self.assertIn(expected, extra)
        self.assertIn("Paragraphs: 9", out)

    def test_the_redlining_validator_accepts_an_unedited_copy(self):
        """``redlining.py`` is what ``validate.py --author`` runs for DOCX."""
        original = self.make_doc(name="original.docx")
        unpacked = self.unpack(original, name="same-unpacked")
        out = self.library(
            f"""
            import json
            from office.validators.redlining import RedliningValidator
            validator = RedliningValidator({str(unpacked)!r}, {str(original)!r})
            print(json.dumps(sorted(
                n for n in dir(validator)
                if not n.startswith("_") and callable(getattr(validator, n)))))
            print(validator.repair())
            print(validator.validate())
            """
        )
        lines = out.splitlines()
        self.assertIn("validate", json.loads(lines[0]))
        self.assertEqual(lines[1], "0", "the redlining validator never repairs")
        self.assertEqual(lines[2], "True")

    def test_the_redlining_validator_rejects_an_untracked_body_edit(self):
        """Ensures the redlining validator rejects an untracked body edit."""
        original = self.make_doc(name="original.docx")
        edited = self.work / "edited.docx"
        document = Document(str(original))
        document.paragraphs[1].text = "Rewritten without tracking"
        self._fix_zoom(document)
        document.save(str(edited))
        unpacked = self.unpack(edited, name="edited-unpacked")

        result = self.import_snippet(
            textwrap.dedent(
                f"""
                from office.validators.redlining import RedliningValidator
                print(RedliningValidator({str(unpacked)!r}, {str(original)!r}).validate())
                """
            ),
            path_entries=[self.scripts],
        )
        self.assertIn("False", result.stdout)
        self.assertAllIn(
            ["FAILED - body text differs from the original",
             "edits made without tracked changes at all",
             "nest <w:del> inside the <w:ins>", "Word-level differences"],
            result.stdout + result.stderr,
        )

    def test_the_pptx_validator_module_backs_the_slides_package(self):
        """The same library shape is shipped for PowerPoint in ``slides``."""
        slides_scripts = ROOT / "skills/slides/scripts"
        out = self.import_snippet(
            textwrap.dedent(
                """
                import json
                from office.validators.base import PackageValidator
                from office.validators.pptx import PPTXSchemaValidator
                print(issubclass(PPTXSchemaValidator, PackageValidator))
                base = {n for n in dir(PackageValidator) if n.startswith("check_")}
                print(json.dumps(sorted(
                    {n for n in dir(PPTXSchemaValidator) if n.startswith("check_")}
                    - base)))
                """
            ),
            path_entries=[slides_scripts],
        ).stdout
        lines = out.splitlines()
        self.assertEqual(lines[0], "True")
        self.assertTrue(json.loads(lines[1]),
                        "the pptx validator must add PowerPoint-specific checks")

    def test_the_soffice_helper_is_shared_by_both_document_packages(self):
        """Ensures the soffice helper is shared by both document packages."""
        for package in ("documents", "slides"):
            with self.subTest(package=package):
                out = self.import_snippet(
                    "import inspect\n"
                    "from office.soffice import run_soffice\n"
                    "print(str(inspect.signature(run_soffice)))\n",
                    path_entries=[ROOT / f"skills/{package}/scripts"],
                ).stdout
                self.assertEqual(out.strip(), "(args, timeout=60, **kwargs)")


@require_modules("pdf2image", "PIL")
class RenderEntrypointTests(DocxFixtureMixin, ScriptCase):
    """``render_docx.py`` is the documented DOCX -> PDF -> PNG entrypoint."""

    def setUp(self):
        super().setUp()
        self.script = ROOT / "skills/documents/render_docx.py"

    def test_help_documents_the_scaling_and_dpi_contract(self):
        """Ensures help documents the scaling and dpi contract."""
        out = self.run_python(self.script, "--help").stdout
        self.assertAllIn(
            ["--output_dir", "--width", "--height", "--dpi", "--emit_pdf", "--verbose",
             "default 1600", "default 2000"],
            out,
        )

    @unittest.skipUnless(binary("soffice") or binary("libreoffice"),
                         "rendering needs LibreOffice")
    def test_renders_png_pages_and_can_emit_the_intermediate_pdf(self):
        """Ensures renders png pages and can emit the intermediate pdf."""
        source = self.make_doc()
        outdir = self.work / "render"
        self.run_python(self.script, source, "--output_dir", outdir, "--dpi", "96",
                        "--emit_pdf", timeout=300)
        pngs = sorted(outdir.glob("*.png"))
        self.assertTrue(pngs, f"no PNGs produced: {sorted(p.name for p in outdir.iterdir())}")
        self.assertTrue(list(outdir.glob("*.pdf")))

    def test_missing_soffice_is_reported_not_silently_ignored(self):
        """Ensures missing soffice is reported not silently ignored."""
        if binary("soffice") or binary("libreoffice"):
            self.skipTest("LibreOffice is installed; the failure path is not reachable")
        result = self.run_python(self.script, self.make_doc(), "--output_dir",
                                 self.work / "none", code=None)
        self.assertNotEqual(result.returncode, 0)


if __name__ == "__main__":
    unittest.main()
